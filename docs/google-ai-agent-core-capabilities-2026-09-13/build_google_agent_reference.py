from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)
DOCX_OUT = ROOT / "Google AI Agent 核心能力参考 2026-09-13.docx"
HTML_OUT = ROOT / "Google AI Agent 核心能力参考 2026-09-13.html"

NAVY = "153B67"
BLUE = "2563EB"
INK = "111827"
MUTED = "475569"
LINE = "D7E1EE"
PALE = "EFF6FF"
PALE_GREEN = "ECFDF5"
GREEN = "047857"
PALE_ORANGE = "FFF7ED"
ORANGE = "C2410C"


def font(size, bold=False):
    candidates = ["C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
                  "C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc"]
    for name in candidates:
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def rounded(draw, box, fill, outline=None, radius=28, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, fnt, fill, spacing=7):
    l, t, r, b = box
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=spacing, align="center")
    draw.multiline_text(((l + r - (bbox[2] - bbox[0])) / 2, (t + b - (bbox[3] - bbox[1])) / 2),
                        text, font=fnt, fill=fill, spacing=spacing, align="center")


def make_capability_map(path):
    im = Image.new("RGB", (1800, 1120), "#F8FBFF")
    d = ImageDraw.Draw(im)
    title, subtitle, big, body, small = font(45, True), font(23), font(34, True), font(23), font(19)
    d.text((80, 54), "Google 官方 AI Agent 能力框架", font=title, fill="#0F2F55")
    d.text((82, 115), "核心能力不是单一模型功能，而是从认知循环到企业运行保障的一整套系统能力。", font=subtitle, fill="#52677F")
    # central loop
    rounded(d, (610, 350, 1190, 715), "#FFFFFF", "#2563EB", 42, 4)
    center_text(d, (650, 385, 1150, 485), "Agent 核心循环", big, "#163B66")
    steps = [(680, 525, 805, 625, "思考\nThink"), (840, 525, 965, 625, "行动\nAct"), (1000, 525, 1125, 625, "观察\nObserve")]
    for x1, y1, x2, y2, text in steps:
        rounded(d, (x1, y1, x2, y2), "#EAF2FF", "#93C5FD", 20, 2)
        center_text(d, (x1, y1, x2, y2), text, body, "#174EA6")
    d.line((805, 575, 840, 575), fill="#2563EB", width=5)
    d.line((965, 575, 1000, 575), fill="#2563EB", width=5)
    d.arc((1050, 610, 1175, 750), start=20, end=160, fill="#2563EB", width=5)
    # surrounding blocks
    blocks = [
        (90, 230, 500, 405, "模型", "推理、规划、生成\n选择适合任务的模型", "#EEF2FF", "#4338CA"),
        (90, 515, 500, 690, "Grounding", "检索、上下文、事实约束\n减少幻觉并引用依据", "#ECFDF5", "#047857"),
        (90, 800, 500, 975, "数据与记忆", "工作记忆、长期知识\n事务记录与可追溯审计", "#FFF7ED", "#C2410C"),
        (1300, 230, 1710, 405, "工具", "API、代码、检索、业务系统\n把答案转为可执行动作", "#EFF6FF", "#1D4ED8"),
        (1300, 515, 1710, 690, "编排", "ReAct 循环、状态、路由\n多 Agent 委派与协同", "#F5F3FF", "#6D28D9"),
        (1300, 800, 1710, 975, "运行时", "安全执行、会话、弹性扩缩\n可靠性、成本与可观测性", "#FFF1F2", "#BE123C"),
    ]
    for x1, y1, x2, y2, head, desc, bg, color in blocks:
        rounded(d, (x1, y1, x2, y2), bg, "#CBD5E1", 28, 2)
        d.text((x1 + 30, y1 + 25), head, font=body, fill=color)
        d.multiline_text((x1 + 30, y1 + 75), desc, font=small, fill="#334155", spacing=8)
    for y in (315, 600, 885):
        d.line((500, y, 610, y + 110 if y != 600 else y), fill="#94A3B8", width=3)
        d.line((1190, y + 110 if y != 600 else y, 1300, y), fill="#94A3B8", width=3)
    rounded(d, (300, 1015, 1500, 1085), "#EAF2FF", "#93C5FD", 20, 2)
    center_text(d, (300, 1015, 1500, 1085), "生产级护栏：身份与权限 · 工具策略 · 沙箱与网络边界 · 轨迹评估 · 日志指标 · 分阶段发布", body, "#153B67")
    im.save(path)


def make_roadmap(path):
    im = Image.new("RGB", (1800, 930), "#FBFDFF")
    d = ImageDraw.Draw(im)
    title, sub, h, body = font(43, True), font(22), font(29, True), font(21)
    d.text((80, 55), "从可用原型到生产级 Agent 的建设路径", font=title, fill="#0F2F55")
    d.text((82, 116), "按风险与收益递进：先让事实可见、动作可控，再沉淀知识并扩展协同。", font=sub, fill="#52677F")
    phases = [
        (85, "P0", "可信问答", "项目上下文与检索\n只读工具\n完整会话与引用记录", "#EFF6FF", "#1D4ED8"),
        (510, "P1", "可控执行", "受控写操作\n审批、回滚与幂等\n工具白名单与审计", "#ECFDF5", "#047857"),
        (935, "P2", "知识与评估", "长期记忆与知识版本\n高价值对话提炼\n轨迹评测与质量基线", "#FFF7ED", "#C2410C"),
        (1360, "P3", "协同与治理", "多 Agent 分工与交接\n身份、策略与注册表\n规模、成本与持续优化", "#F5F3FF", "#6D28D9"),
    ]
    for x, tag, head, desc, bg, color in phases:
        rounded(d, (x, 245, x + 355, 685), bg, "#CBD5E1", 35, 2)
        rounded(d, (x + 28, 275, x + 110, 330), color, color, 18, 1)
        center_text(d, (x + 28, 275, x + 110, 330), tag, font(21, True), "#FFFFFF")
        d.text((x + 28, 370), head, font=h, fill="#172554")
        d.multiline_text((x + 28, 435), desc, font=body, fill="#334155", spacing=15)
        if x < 1360:
            d.line((x + 355, 465, x + 405, 465), fill="#94A3B8", width=5)
            d.polygon([(x + 405, 465), (x + 387, 452), (x + 387, 478)], fill="#94A3B8")
    d.line((145, 760, 1655, 760), fill="#CBD5E1", width=2)
    center_text(d, (200, 785, 1600, 860), "每一阶段均持续保留：业务目标、事实来源、权限边界、执行证据、失败复盘和质量指标。", body, "#334155")
    im.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=LINE):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def set_cell_margin(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def configure(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.72), Inches(0.7)
    sec.left_margin, sec.right_margin = Inches(0.72), Inches(0.72)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size, normal.font.color.rgb = Pt(10.5), RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.32
    for key, size, color in (("Title", 26, "000000"), ("Heading 1", 16, "000000"), ("Heading 2", 12.5, "000000")):
        st = styles[key]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size, st.font.bold, st.font.color.rgb = Pt(size), True, RGBColor.from_string(color)
    styles["Heading 1"].paragraph_format.space_before, styles["Heading 1"].paragraph_format.space_after = Pt(18), Pt(8)
    styles["Heading 2"].paragraph_format.space_before, styles["Heading 2"].paragraph_format.space_after = Pt(12), Pt(5)


def add_text(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margin(cell)
        run = cell.paragraphs[0].add_run(value)
        run.bold, run.font.color.rgb, run.font.size = True, RGBColor(255, 255, 255), Pt(9.5)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F7FAFE" if index % 2 == 0 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margin(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, path, caption, width=7.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor.from_string(MUTED)
    cap.paragraph_format.space_after = Pt(11)


def set_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Google AI Agent 核心能力参考  |  2026-09-13")
    r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string("64748B")


def build_docx():
    map_file = ASSETS / "google-agent-capability-map.png"
    roadmap_file = ASSETS / "google-agent-production-path.png"
    make_capability_map(map_file)
    make_roadmap(roadmap_file)
    doc = Document()
    configure(doc)
    set_footer(doc.sections[0])
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Google AI Agent 核心能力参考")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向产品规划、技术建设与上线治理的官方材料整理  |  2026-09-13")
    r.font.size, r.font.color.rgb = Pt(11), RGBColor.from_string(MUTED)
    doc.add_paragraph()
    add_text(doc, "结论：Google Cloud 的官方材料把 AI Agent 视为一套完整系统，而不只是“能聊天的模型”。核心由模型、Grounding、工具、数据架构、编排和运行时组成；要在真实业务中稳定运行，还必须同时建设安全与身份、评估与可观测性、可靠交付与成本治理。")
    add_text(doc, "本文将 Google 官方页面中分散的能力描述归并为一套便于规划和验收的清单。它是对官方材料的综合整理，不等同于某一个 Google 产品或 API 的单独功能列表。")

    doc.add_heading("1. Google 官方框架中的六个构建块", level=1)
    add_text(doc, "Google Cloud 在 Core concepts of AI agents 中明确列出六个构建块。它们共同决定 Agent 是否能理解目标、获得可靠上下文、采取行动并在生产环境持续运行。")
    add_table(doc, ["构建块", "解决的问题", "建设要点"], [
        ("模型", "理解任务、推理、生成计划与结果", "按任务选择模型；定义系统指令、输出结构和失败降级策略。"),
        ("Grounding", "把回答约束在可信事实和当前上下文中", "检索项目文档、代码、业务数据；要求来源、时间范围和权限边界可见。"),
        ("工具", "让 Agent 调用外部系统而不是只给建议", "以明确输入输出封装 API、Git、检索、工单等动作；设定批准、回滚和审计。"),
        ("数据架构", "保存知识、记忆、状态与执行证据", "区分工作记忆、长期记忆、业务数据和事务日志；治理版本、权限与保留期。"),
        ("编排", "组织多轮循环、状态流转和多 Agent 协作", "实现任务拆分、路由、委派、重试、交接与停止条件。"),
        ("运行时", "在真实环境安全、可靠、可扩展地执行", "提供隔离、身份、网络控制、弹性、观测、评估和成本管理。"),
    ], [1.05, 2.25, 3.72])
    add_figure(doc, map_file, "图 1  Google 官方材料归纳的 Agent 能力全景图")

    doc.add_heading("2. Agent 的核心闭环", level=1)
    add_text(doc, "Google 将生产级 Agent 的工作方式描述为递归的 Think - Act - Observe 循环，并由编排层维护状态与下一步决策。对业务系统而言，闭环是否完整，比单次回答是否流畅更重要。")
    add_table(doc, ["环节", "应具备的能力", "验收信号"], [
        ("思考 Think", "理解目标、识别约束、拆分步骤、决定是否要提问或调用工具。", "能解释为什么选择某条路径；对不确定信息明确说明。"),
        ("行动 Act", "按受控接口调用检索、代码、业务 API 或委派其他 Agent。", "每次动作带有输入、权限、结果和关联会话。"),
        ("观察 Observe", "读取工具返回、识别异常、修正计划，必要时恢复或停止。", "错误可定位；支持重试、换路、请求澄清和安全停止。"),
        ("迭代与交付", "更新短期状态，生成结果、依据和后续建议。", "输出可追溯到事实、工具结果和执行轨迹。"),
    ], [1.1, 3.55, 2.37])

    doc.add_heading("3. Grounding、上下文与记忆", level=1)
    add_text(doc, "Google 的材料强调，Grounding 不只是给模型追加一段 RAG 文本。一个生产级 Agent 需要把实时上下文、可复用知识和动作证据分层处理，避免把不可靠或无权限的信息混入决策。")
    add_table(doc, ["数据层", "用途", "建议的治理方式"], [
        ("工作记忆", "当前会话目标、已读文件、已调用工具、未完成步骤。", "随任务保存；限制大小；切换会话时恢复位置与关键状态。"),
        ("长期知识与记忆", "项目规范、架构、历史决策、用户偏好和经过提炼的高价值经验。", "来源可追溯、版本化、可撤销；区分项目、团队和个人可见范围。"),
        ("事务与审计记录", "写操作、Git 提交、工单更新、审批、失败和补偿过程。", "不可随意覆盖；关联执行者、时间、输入、输出、权限和结果。"),
    ], [1.25, 3.1, 2.67])
    add_text(doc, "实践提示：对代码与交付场景，建议将“代码事实”“需求事实”“执行证据”分开保存。模型可以综合它们回答，但不能把推测写回事实库。")

    doc.add_heading("4. 工具调用与受控执行", level=1)
    add_text(doc, "工具让 Agent 从问答走向工作流，但也是风险最大的边界。Google 的生产实践强调围绕工具调用构建安全、可恢复和可观察的执行机制。")
    add_bullets(doc, [
        "把每个动作封装为边界清晰的工具：输入、输出、权限、超时、幂等性和失败语义均应明确。",
        "写入、删除、推送、发布等高影响动作采用分级批准；只读查询可自动执行，危险动作应提供预览与确认。",
        "用最小权限身份执行，并限制网络出口、文件范围和可调用工具；敏感信息不应直接出现在模型上下文或日志中。",
        "对失败设计重试、去重、补偿与人工接管；不要把“模型再次尝试”当作唯一恢复策略。",
    ])

    doc.add_heading("5. 编排与多 Agent 协同", level=1)
    add_text(doc, "Google 将编排视为把模型、工具和数据连接为可运行工作流的控制层。多 Agent 的价值不是增加角色数量，而是让不同职责在统一状态、权限和交接规范下完成分工。")
    add_table(doc, ["协同能力", "适用场景", "应保留的控制点"], [
        ("任务路由", "根据项目、专业、风险或工具范围选择合适的 Agent。", "能力声明、模型与工具范围、超时与失败回退。"),
        ("委派与交接", "分析 Agent 将实现任务交给代码 Agent，或把待确认项交给人工。", "输入摘要、事实来源、未完成项、责任归属与交接结果。"),
        ("共享状态", "多个 Agent 围绕同一任务持续推进。", "区分共享任务状态和私有推理；并发写入需要锁、版本或事务策略。"),
        ("互操作", "将外部 Agent、工具服务器或业务系统接入。", "标准协议只是接口；仍要做身份验证、权限、观测和内容安全。"),
    ], [1.25, 3.15, 2.62])

    doc.add_heading("6. 生产级能力：安全、评估与可观测性", level=1)
    add_text(doc, "Google Cloud 的生产实践不把安全和评估视为上线后的附加项，而是 Agent 运行时的一部分。评估对象也不应只看最终答案，还应查看完整轨迹。")
    add_table(doc, ["领域", "关键问题", "建议指标或证据"], [
        ("身份与策略", "谁在代表谁执行？能访问什么？", "Agent 身份、最小权限、工具白名单、策略命中记录、审批记录。"),
        ("内容与数据安全", "是否被提示注入、工具投毒或敏感数据泄露影响？", "输入输出过滤、数据分级、密钥隔离、异常拦截与人工复核。"),
        ("轨迹评估", "是否选择了正确工具、合理推理并从错误中恢复？", "工具选择正确率、澄清率、失败恢复率、人工接管率、任务完成质量。"),
        ("可观测性", "发生了什么，为什么发生，成本如何？", "Trace、模型调用、工具耗时、错误类别、上下文大小、Token 和单位任务成本。"),
        ("发布与可靠性", "新版本是否在真实业务中安全生效？", "离线评测、沙箱、灰度、回滚、SLO、限流与容量指标。"),
    ], [1.25, 3.05, 2.72])

    doc.add_heading("7. 建设顺序建议", level=1)
    add_text(doc, "以下路径是基于 Google 官方能力框架作出的工程规划建议，不是 Google 对任何具体产品的实施要求。它适用于希望从代码问答逐步扩展到可协作交付的团队。")
    add_figure(doc, roadmap_file, "图 2  从可信问答到多 Agent 治理的递进建设路径")
    add_table(doc, ["阶段", "优先交付", "暂不宜过早追求"], [
        ("P0 可信问答", "项目选择、事实检索、只读工具、引用与会话轨迹。", "无人值守写库、复杂多 Agent 自主决策。"),
        ("P1 可控执行", "受控代码修改、Git 证据链、审批、回滚、幂等与错误处理。", "让模型拥有广泛系统权限。"),
        ("P2 知识与评估", "高价值对话提炼、项目知识版本、评测集、轨迹分析和质量基线。", "把所有原始对话自动写入长期记忆。"),
        ("P3 协同与治理", "专业 Agent、委派交接、统一身份策略、Agent 注册与成本优化。", "无策略边界的自治协作。"),
    ], [1.28, 3.25, 2.49])

    doc.add_heading("8. 对坤伴后续规划的映射", level=1)
    add_text(doc, "从已有的项目、会话、工作画布、代码库、Git 操作和多模型接入能力出发，可将后续建设聚焦在“可靠事实 + 可控行动 + 可沉淀知识 + 可协同交付”四条主线。以下为规划推断，非 Google 官方对坤伴的评价。")
    add_bullets(doc, [
        "统一知识入口：以项目、代码、需求、任务和交付记录形成可引用的知识对象；为来源、版本、权限和有效期建立元数据。",
        "可追溯记忆：从高价值对话中提炼决策、约束、故障经验和技能，但先经规则或人工确认，再进入长期知识。",
        "证据链执行：把工具调用、文件差异、Git 提交、审批、测试与发布结果串成一条可回放的任务轨迹。",
        "Agent 协同：先定义能力、输入输出、权限与交接格式，再扩展 CPS、APS 等专业 Agent 的协作，而不是只增加聊天入口。",
        "质量运营：建设典型任务集和回归评测，持续观察正确性、失败恢复、成本、时延与人工介入比例。",
    ])

    doc.add_heading("附录 A 主要 Google 官方来源", level=1)
    sources = [
        ("Google Cloud - Core concepts of AI agents", "https://cloud.google.com/resources/core-concepts-ai-agents"),
        ("Google Cloud - What are AI agents", "https://cloud.google.com/discover/what-are-ai-agents"),
        ("Google Cloud Blog - A developer's guide to production-ready AI agents", "https://cloud.google.com/blog/products/ai-machine-learning/a-devs-guide-to-production-ready-ai-agents"),
        ("Google Cloud - What are agentic workflows", "https://cloud.google.com/discover/agentic-workflows"),
        ("Google Cloud documentation - Host AI agents on Cloud Run", "https://docs.cloud.google.com/run/docs/ai-agents"),
        ("Google AI for Developers - Agents overview", "https://ai.google.dev/gemini-api/docs/agents"),
    ]
    for name, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(name + "\n")
        r.bold = True
        u = p.add_run(url)
        u.font.size, u.font.color.rgb = Pt(9), RGBColor.from_string(BLUE)
    doc.save(DOCX_OUT)


def build_html():
    html = '''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Google AI Agent 核心能力参考</title>
<style>
:root{--navy:#153b67;--blue:#2563eb;--ink:#111827;--muted:#475569;--line:#d7e1ee;--bg:#f7faff}*{box-sizing:border-box}body{margin:0;color:var(--ink);background:var(--bg);font:16px/1.7 "Microsoft YaHei",Arial,sans-serif}.top{background:linear-gradient(120deg,#102c4e,#2563eb);color:#fff;padding:60px max(5vw,30px) 52px}.top h1{margin:0;font-size:clamp(28px,4vw,48px)}.top p{max-width:850px;margin:12px 0 0;color:#dbeafe}.wrap{max-width:1180px;margin:auto;padding:32px 24px 64px}.layout{display:grid;grid-template-columns:230px minmax(0,1fr);gap:28px}.nav{position:sticky;top:20px;height:max-content;background:#fff;border:1px solid var(--line);border-radius:16px;padding:14px}.nav a{display:block;padding:8px 10px;text-decoration:none;color:var(--muted);border-radius:8px}.nav a:hover{background:#eff6ff;color:var(--blue)}main{min-width:0}section{background:#fff;border:1px solid var(--line);border-radius:18px;padding:28px;margin-bottom:24px;box-shadow:0 10px 28px #153b6709}h2{margin:0 0 10px;color:var(--navy);font-size:24px}h3{color:#173b67;margin:22px 0 8px}.lead{font-size:18px}.figure{width:100%;border:1px solid var(--line);border-radius:14px;margin:18px 0 5px}.caption{text-align:center;color:var(--muted);font-size:13px;margin:0 0 18px}table{border-collapse:separate;border-spacing:0;width:100%;overflow:hidden;border:1px solid var(--line);border-radius:10px;margin:14px 0}th{background:var(--navy);color:#fff;text-align:left}th,td{padding:12px 13px;border-right:1px solid var(--line);border-bottom:1px solid var(--line);vertical-align:top}td:last-child,th:last-child{border-right:0}tr:last-child td{border-bottom:0}tbody tr:nth-child(odd){background:#f8fbff}.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}.card{border:1px solid var(--line);border-radius:12px;padding:16px}.card b{color:var(--blue)}details{border:1px solid var(--line);border-radius:10px;padding:12px 14px;margin:10px 0}summary{cursor:pointer;font-weight:700;color:var(--navy)}a{color:#1d4ed8}li{margin:6px 0}.tag{display:inline-block;border-radius:999px;background:#dbeafe;color:#174ea6;padding:3px 9px;font-size:12px;font-weight:700;margin:0 6px 6px 0}.foot{color:var(--muted);font-size:13px}@media(max-width:800px){.layout{grid-template-columns:1fr}.nav{display:none}.wrap{padding:20px 14px}.top{padding:42px 22px}section{padding:20px;overflow:auto}.grid{grid-template-columns:1fr}table{min-width:650px}}
</style></head><body><header class="top"><h1>Google AI Agent 核心能力参考</h1><p>面向产品规划、技术建设与上线治理的官方材料整理 · 2026-09-13</p></header><div class="wrap"><div class="layout"><nav class="nav"><a href="#six">六个构建块</a><a href="#loop">核心闭环</a><a href="#memory">上下文与记忆</a><a href="#tools">受控执行</a><a href="#multi">多 Agent 协同</a><a href="#production">生产级能力</a><a href="#roadmap">建设路径</a><a href="#sources">官方来源</a></nav><main>
<section><p class="lead"><b>结论：</b>Google Cloud 的官方材料把 AI Agent 视为一套完整系统，而不只是“能聊天的模型”。核心由模型、Grounding、工具、数据架构、编排和运行时组成；上线还必须同步建设安全、评估、观测、可靠交付与成本治理。</p><p>本文是对 Google 官方材料的综合整理，不等同于某一个 Google 产品或 API 的单独功能清单。</p></section>
<section id="six"><h2>1. 六个构建块</h2><p>Google Cloud 在 Core concepts of AI agents 中明确列出以下六项。</p><img class="figure" src="assets/google-agent-capability-map.png" alt="Google AI Agent 核心能力框架图"><p class="caption">图 1 · 核心能力全景图</p><table><thead><tr><th>构建块</th><th>解决的问题</th><th>建设要点</th></tr></thead><tbody><tr><td>模型</td><td>理解、推理、计划、生成</td><td>模型选择、系统指令、结构化输出和失败降级。</td></tr><tr><td>Grounding</td><td>可信事实与当前上下文</td><td>检索项目文档、代码和业务数据，显示来源与权限。</td></tr><tr><td>工具</td><td>将答案转为可执行动作</td><td>封装 API、Git、检索、工单等动作，配置审批和审计。</td></tr><tr><td>数据架构</td><td>知识、记忆、状态与证据</td><td>分层管理工作记忆、长期知识、业务数据和事务日志。</td></tr><tr><td>编排</td><td>多轮循环与协作</td><td>任务拆分、路由、委派、重试、交接与停止条件。</td></tr><tr><td>运行时</td><td>安全可靠地运行</td><td>隔离、身份、网络控制、弹性、观测、评估和成本管理。</td></tr></tbody></table></section>
<section id="loop"><h2>2. 核心闭环</h2><p>生产级 Agent 的行为可归纳为递归的 <b>Think → Act → Observe</b> 循环，由编排层保存状态并决定下一步。</p><div class="grid"><div class="card"><b>思考 Think</b><br>理解目标、识别约束、拆分步骤，判断是否需要澄清或调用工具。</div><div class="card"><b>行动 Act</b><br>通过受控接口调用检索、代码、业务 API 或委派其他 Agent。</div><div class="card"><b>观察 Observe</b><br>读取工具返回、识别异常、修正计划，必要时恢复或停止。</div><div class="card"><b>交付与迭代</b><br>更新状态，输出结果、依据和后续建议，保留可追溯轨迹。</div></div></section>
<section id="memory"><h2>3. Grounding、上下文与记忆</h2><p>Grounding 不只是 RAG 文本。生产 Agent 需要分层处理实时上下文、可复用知识与动作证据。</p><details open><summary>工作记忆</summary>保存当前会话目标、已读文件、已调用工具和未完成步骤。应限制大小，并支持恢复任务位置。</details><details><summary>长期知识与记忆</summary>沉淀项目规范、架构、历史决策、用户偏好和经确认的高价值经验。必须可追溯、版本化、可撤销，并区分项目、团队和个人权限。</details><details><summary>事务与审计记录</summary>保存写操作、Git 提交、工单更新、审批、失败与补偿。关联执行者、时间、输入、输出、权限和结果。</details><p><b>实践提示：</b>代码事实、需求事实和执行证据应该分开保存；模型可综合它们回答，但不能把推测写回事实库。</p></section>
<section id="tools"><h2>4. 工具调用与受控执行</h2><ul><li>每个工具定义输入、输出、权限、超时、幂等性和失败语义。</li><li>写入、删除、推送和发布等高影响动作采用分级批准，提供预览和确认。</li><li>使用最小权限身份，限制网络出口、文件范围和可调用工具。</li><li>对失败设计重试、去重、补偿与人工接管，不能只依赖模型再次尝试。</li></ul></section>
<section id="multi"><h2>5. 多 Agent 协同</h2><p>多 Agent 的价值在于专业分工与受控交接，不在于角色数量。</p><table><thead><tr><th>能力</th><th>用途</th><th>控制点</th></tr></thead><tbody><tr><td>任务路由</td><td>按项目、专业、风险选择 Agent</td><td>能力声明、模型与工具范围、超时、失败回退</td></tr><tr><td>委派与交接</td><td>分析、实现、测试和人工确认分工</td><td>输入摘要、事实来源、未完成项、责任归属</td></tr><tr><td>共享状态</td><td>多个 Agent 围绕同一任务推进</td><td>共享任务状态与私有推理分离；并发写入可控</td></tr><tr><td>互操作</td><td>连接外部 Agent、工具服务器和业务系统</td><td>标准协议之外仍需身份、权限、观测和内容安全</td></tr></tbody></table></section>
<section id="production"><h2>6. 生产级能力</h2><table><thead><tr><th>领域</th><th>关键问题</th><th>证据或指标</th></tr></thead><tbody><tr><td>身份与策略</td><td>谁在代表谁执行？能访问什么？</td><td>Agent 身份、最小权限、白名单、审批记录</td></tr><tr><td>内容与数据安全</td><td>是否遭遇提示注入、工具投毒或泄露？</td><td>输入输出过滤、数据分级、密钥隔离、异常拦截</td></tr><tr><td>轨迹评估</td><td>工具和推理路径是否合理？</td><td>工具选择、澄清、恢复、人工接管和完成质量</td></tr><tr><td>可观测性</td><td>发生了什么，为什么，成本如何？</td><td>Trace、耗时、错误、上下文大小、Token、单位成本</td></tr><tr><td>发布与可靠性</td><td>新版本如何安全进入生产？</td><td>离线评测、沙箱、灰度、回滚、SLO、限流</td></tr></tbody></table></section>
<section id="roadmap"><h2>7. 建设路径建议</h2><p>以下是基于 Google 官方能力框架形成的工程规划建议，不是 Google 对具体产品的实施要求。</p><img class="figure" src="assets/google-agent-production-path.png" alt="从原型到生产的 AI Agent 建设路径"><p class="caption">图 2 · 从可信问答到多 Agent 治理</p><p><span class="tag">P0 可信问答</span><span class="tag">P1 可控执行</span><span class="tag">P2 知识与评估</span><span class="tag">P3 协同与治理</span></p><ul><li><b>P0：</b>项目上下文、事实检索、只读工具、引用与会话轨迹。</li><li><b>P1：</b>受控修改、Git 证据链、审批、回滚、幂等和错误处理。</li><li><b>P2：</b>高价值对话提炼、知识版本、评测集、轨迹分析与质量基线。</li><li><b>P3：</b>专业 Agent、委派交接、统一身份策略、Agent 注册与成本优化。</li></ul></section>
<section><h2>8. 对坤伴后续规划的映射</h2><p>从项目、会话、工作画布、代码库、Git 操作和多模型接入能力出发，建议聚焦“可靠事实、可控行动、可沉淀知识、可协同交付”。这是规划推断，非 Google 官方评价。</p><ul><li>建设有来源、版本、权限和有效期的知识对象。</li><li>提炼经确认的高价值对话为长期记忆，而非无差别保存。</li><li>把工具调用、文件差异、Git 提交、审批、测试和发布串为可回放证据链。</li><li>先定义能力、输入输出、权限与交接格式，再扩展 CPS、APS 等专业 Agent。</li><li>用典型任务集持续观察正确性、失败恢复、成本、时延与人工介入比例。</li></ul></section>
<section id="sources"><h2>附录 · Google 官方来源</h2><ul><li><a href="https://cloud.google.com/resources/core-concepts-ai-agents">Google Cloud · Core concepts of AI agents</a></li><li><a href="https://cloud.google.com/discover/what-are-ai-agents">Google Cloud · What are AI agents</a></li><li><a href="https://cloud.google.com/blog/products/ai-machine-learning/a-devs-guide-to-production-ready-ai-agents">Google Cloud Blog · A developer's guide to production-ready AI agents</a></li><li><a href="https://cloud.google.com/discover/agentic-workflows">Google Cloud · What are agentic workflows</a></li><li><a href="https://docs.cloud.google.com/run/docs/ai-agents">Google Cloud documentation · Host AI agents on Cloud Run</a></li><li><a href="https://ai.google.dev/gemini-api/docs/agents">Google AI for Developers · Agents overview</a></li></ul><p class="foot">整理日期：2026-09-13。Google 页面会随产品和文档更新而变化；实际落地应以目标产品的最新文档、权限和合规要求为准。</p></section>
</main></div></div></body></html>'''
    HTML_OUT.write_text(html, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_html()
    print(DOCX_OUT)
    print(HTML_OUT)
