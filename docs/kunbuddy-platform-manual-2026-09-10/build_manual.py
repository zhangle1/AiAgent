from __future__ import annotations

import argparse
import io
import json
import math
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def font(size: int, bold: bool = False):
    names = [
        "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc",
    ]
    for name in names:
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def extract_source_images(source: Path, assets: Path) -> list[Path]:
    target = assets / "source-screenshots"
    target.mkdir(parents=True, exist_ok=True)
    images: list[Path] = []
    with zipfile.ZipFile(source) as archive:
        names = sorted(
            (name for name in archive.namelist() if name.startswith("word/media/image")),
            key=lambda item: int(item.rsplit("image", 1)[1].split(".", 1)[0]),
        )
        for index, name in enumerate(names, start=1):
            extension = Path(name).suffix.lower() or ".png"
            output = target / f"screen-{index:02d}{extension}"
            output.write_bytes(archive.read(name))
            images.append(output)
    return images


def make_contact_sheet(images: list[Path], output: Path) -> None:
    cell_width, cell_height = 330, 230
    cols = 3
    rows = math.ceil(len(images) / cols)
    canvas = Image.new("RGB", (cols * cell_width + 40, rows * cell_height + 40), "#f8fafc")
    draw = ImageDraw.Draw(canvas)
    label_font = font(16, True)
    for idx, path in enumerate(images):
        with Image.open(path) as raw:
            image = raw.convert("RGB")
        image.thumbnail((cell_width - 24, cell_height - 52), Image.Resampling.LANCZOS)
        x = 20 + (idx % cols) * cell_width
        y = 20 + (idx // cols) * cell_height
        draw.rounded_rectangle((x, y, x + cell_width - 12, y + cell_height - 12), radius=12, fill="white", outline="#cbd5e1")
        draw.text((x + 12, y + 10), f"{idx + 1:02d}  {path.name}", font=label_font, fill="#0f172a")
        image_x = x + (cell_width - 12 - image.width) // 2
        image_y = y + 38 + (cell_height - 50 - image.height) // 2
        canvas.paste(image, (image_x, image_y))
    canvas.save(output)


def centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, text_font, color: str, spacing: int = 8) -> None:
    """Draw a multi-line label centered in a rectangular region."""
    left, top, right, bottom = box
    bounds = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=spacing, align="center")
    x = left + (right - left - (bounds[2] - bounds[0])) / 2
    y = top + (bottom - top - (bounds[3] - bounds[1])) / 2
    draw.multiline_text((x, y), text, font=text_font, fill=color, spacing=spacing, align="center")


def diagram_card(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    fill: str = "#ffffff",
    accent: str = "#2563eb",
) -> None:
    left, top, right, bottom = box
    draw.rounded_rectangle(box, radius=26, fill=fill, outline="#cbd5e1", width=3)
    draw.rounded_rectangle((left, top, left + 14, bottom), radius=7, fill=accent)
    title_font = font(31, True)
    body_font = font(23)
    draw.text((left + 36, top + 28), title, font=title_font, fill="#14213d")
    draw.multiline_text((left + 36, top + 80), body, font=body_font, fill="#475569", spacing=10)


def diagram_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#4f7ef7") -> None:
    draw.line((start, end), fill=color, width=8)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        draw.polygon([(x2, y2), (x2 - 24 * direction, y2 - 14), (x2 - 24 * direction, y2 + 14)], fill=color)
    else:
        direction = 1 if y2 > y1 else -1
        draw.polygon([(x2, y2), (x2 - 14, y2 - 24 * direction), (x2 + 14, y2 - 24 * direction)], fill=color)


def diagram_elbow_arrow(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]], color: str) -> None:
    """Route a connector around cards and finish it with a directional arrow."""
    draw.line(points, fill=color, width=8, joint="curve")
    start = points[-2]
    end = points[-1]
    x1, y1 = start
    x2, y2 = end
    direction = 1 if x2 > x1 else -1
    if abs(x2 - x1) >= abs(y2 - y1):
        draw.polygon([(x2, y2), (x2 - 24 * direction, y2 - 14), (x2 - 24 * direction, y2 + 14)], fill=color)
    else:
        direction = 1 if y2 > y1 else -1
        draw.polygon([(x2, y2), (x2 - 14, y2 - 24 * direction), (x2 + 14, y2 - 24 * direction)], fill=color)


def build_knowledge_diagrams(assets: Path) -> dict[str, Path]:
    """Generate two original planning diagrams for the knowledge and Agent roadmap."""
    assets.mkdir(parents=True, exist_ok=True)
    title_font = font(42, True)
    sub_font = font(24)
    small_font = font(21)

    flow = Image.new("RGB", (1800, 1080), "#f8fafc")
    draw = ImageDraw.Draw(flow)
    draw.rounded_rectangle((30, 28, 1770, 1050), radius=34, fill="#ffffff", outline="#dbe5f2", width=3)
    draw.text((80, 68), "坤伴知识库与 CPS APS Agent 协同流程", font=title_font, fill="#14213d")
    draw.text((82, 126), "以受控项目事实为入口，按需取数、跨 Agent 复用，并把结论与交付证据沉淀回知识资产。", font=sub_font, fill="#64748b")

    diagram_card(draw, (90, 235, 475, 430), "CPS 项目事实", "代码仓库  •  Git 提交\n接口文档  •  配置  •  日志", fill="#f7fbff", accent="#2563eb")
    diagram_card(draw, (90, 505, 475, 700), "APS 项目事实", "计划模型  •  业务规则\n任务附件  •  需求变更", fill="#f7fbff", accent="#7c3aed")
    diagram_card(draw, (585, 300, 1050, 655), "知识资产中心", "资源分层：摘要  /  概览  /  明细\n可追溯 URI：kun://project/...\n版本、权限、来源、有效期\n检索轨迹与引用证据", fill="#eff6ff", accent="#0891b2")
    diagram_card(draw, (1165, 220, 1650, 430), "CPS Agent", "定位 Bug 与代码差异\n分析接口影响  •  生成变更集\n以仓库版本作为事实依据", fill="#f0fdf4", accent="#16a34a")
    diagram_card(draw, (1165, 520, 1650, 730), "APS Agent", "解释计划逻辑与约束\n评估需求影响  •  输出方案\n引用任务与规则证据", fill="#faf5ff", accent="#7c3aed")
    diagram_card(draw, (585, 785, 1050, 955), "协同与治理", "开发 / 交付共同复核  •  变更集审批\nGit 提交、测试结果、检索轨迹形成证据链", fill="#fffaf0", accent="#d97706")

    diagram_arrow(draw, (475, 332), (585, 405))
    diagram_arrow(draw, (475, 602), (585, 550))
    diagram_arrow(draw, (1050, 405), (1165, 325), "#16a34a")
    diagram_arrow(draw, (1050, 550), (1165, 625), "#7c3aed")
    diagram_elbow_arrow(draw, [(1650, 325), (1710, 325), (1710, 760), (1000, 760), (1000, 785)], "#0f766e")
    diagram_elbow_arrow(draw, [(1650, 625), (1680, 625), (1680, 805), (1030, 805), (1030, 785)], "#7c3aed")
    centered_text(draw, (105, 880, 490, 975), "采集时保留来源、版本和授权范围\n不把外部附件直接当作可信事实", small_font, "#64748b")
    flow_path = assets / "knowledge-cps-aps-collaboration.png"
    flow.save(flow_path)

    roadmap = Image.new("RGB", (1800, 1020), "#f8fafc")
    draw = ImageDraw.Draw(roadmap)
    draw.rounded_rectangle((30, 28, 1770, 990), radius=34, fill="#ffffff", outline="#dbe5f2", width=3)
    draw.text((80, 68), "坤伴 2.0 知识库建设路径", font=title_font, fill="#14213d")
    draw.text((82, 126), "建议按“可控接入 → 可用知识 → Agent 协同 → 自我演进”逐步建设，先确保事实、权限和评估，再扩大自动化范围。", font=sub_font, fill="#64748b")
    phases = [
        ("P0 数据接入", "项目、仓库、文档、附件\n记录来源、版本、权限\n建立增量同步和失效机制", "#2563eb"),
        ("P1 知识可用", "L0 摘要 / L1 概览 / L2 明细\nkun URI 与可追溯引用\n检索、引用、回收与纠错", "#0891b2"),
        ("P2 Agent 协同", "CPS 与 APS Adapter\n共享检索、会话记忆与任务上下文\n复核跨系统影响", "#7c3aed"),
        ("P3 自我演进", "高价值对话提炼为候选知识\n技能评估、版本发布、灰度回滚\n指标驱动的持续优化", "#16a34a"),
    ]
    x_positions = [82, 512, 942, 1372]
    for index, ((phase, body, color), x) in enumerate(zip(phases, x_positions), start=1):
        draw.rounded_rectangle((x, 255, x + 350, 650), radius=28, fill="#ffffff", outline="#cbd5e1", width=3)
        draw.rounded_rectangle((x, 255, x + 350, 330), radius=28, fill=color)
        draw.rectangle((x, 300, x + 350, 330), fill=color)
        centered_text(draw, (x + 20, 267, x + 330, 317), phase, font(29, True), "#ffffff")
        draw.multiline_text((x + 32, 365), body, font=font(24), fill="#475569", spacing=14)
        draw.ellipse((x + 30, 560, x + 68, 598), fill="#eef5ff", outline=color, width=2)
        draw.text((x + 82, 563), f"阶段 {index}", font=small_font, fill=color)
        if index < len(phases):
            diagram_arrow(draw, (x + 350, 452), (x + 420, 452), "#94a3b8")
    draw.rounded_rectangle((82, 755, 1720, 915), radius=24, fill="#f8fbff", outline="#bfdbfe", width=2)
    draw.text((116, 786), "全程治理要求", font=font(28, True), fill="#173b69")
    draw.text((116, 842), "项目级授权  •  事实来源与版本  •  检索和工具调用轨迹  •  人工审批  •  质量评估  •  敏感数据分级与脱敏", font=font(25), fill="#334155")
    roadmap_path = assets / "knowledge-2.0-roadmap.png"
    roadmap.save(roadmap_path)
    return {"knowledge_flow": flow_path, "knowledge_roadmap": roadmap_path}


def inspect(source: Path, output: Path) -> None:
    assets = output / "assets"
    images = extract_source_images(source, assets)
    preview = output / "source-screens-contact-sheet.png"
    make_contact_sheet(images, preview)
    print(json.dumps({"image_count": len(images), "preview": str(preview), "images": [str(item) for item in images]}, ensure_ascii=False, indent=2))


NAVY = "173B69"
BLUE = "2563EB"
PALE_BLUE = "EEF5FF"
LIGHT_BORDER = "D9E2F0"
TEXT = "1F2937"
MUTED = "64748B"


def set_font(run, size: float, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_fill(cell, value: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), value)


def set_cell_borders(cell, color: str = LIGHT_BORDER) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margin(cell, top: int = 120, start: int = 130, bottom: int = 120, end: int = 130) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_table_width(table, widths: list[float]) -> None:
    # Word otherwise expands a Table Grid to the page edge, which can clip the
    # rightmost column in exported PDFs. Set an explicit preferred table width.
    properties = table._tbl.tblPr
    preferred_width = properties.first_child_found_in("w:tblW")
    if preferred_width is None:
        preferred_width = OxmlElement("w:tblW")
        properties.append(preferred_width)
    preferred_width.set(qn("w:w"), str(round(sum(widths) * 1440)))
    preferred_width.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def paragraph(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.42
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_font(run, 11, True)
        remainder = p.add_run(text[len(bold_prefix):])
        set_font(remainder, 11)
    else:
        run = p.add_run(text)
        set_font(run, 11)
    return p


def heading(doc: Document, text: str, level: int) -> None:
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 13, True, "000000")


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_font(run, 10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_table_width(table, widths)
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for cell, text in zip(header_cells, headers):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margin(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_font(run, 10, True, "FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            set_cell_margin(cell)
            if row_index % 2 == 1:
                set_cell_fill(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            run = p.add_run(text)
            set_font(run, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_figure(doc: Document, image_path: Path, caption: str, width: float = 6.75) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    run = cap.add_run(caption)
    set_font(run, 9.5, False, MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    # Keep the text block generous while allowing the 7.1-inch reference tables
    # to remain inside printable page bounds.
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for name, size in (("Title", 27), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor(0, 0, 0)
    # Avoid the decorative bottom rule that some Word installations apply to Title.
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(3)
    run = footer.add_run("坤伴平台操作手册  |  2026-09-10")
    set_font(run, 8.5, False, MUTED)


def add_cover(doc: Document, images: dict[int, Path]) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("坤伴平台操作手册")
    set_font(run, 27, True, "000000")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("面向研发协作、代码交付与平台管理")
    set_font(run, 13, False, MUTED)
    add_figure(doc, images[1], "图 1  坤伴平台当前系统架构", width=6.55)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(20)
    for line in ("版本：1.0 操作与管理指南", "更新日期：2026-09-10", "适用对象：开发人员、项目负责人、代码审批人和系统管理员"):
        run = info.add_run(line + "\n")
        set_font(run, 10.5, False, MUTED)
    doc.add_page_break()


def build_docx(output: Path, images: dict[int, Path], login_image: Path, knowledge_diagrams: dict[str, Path]) -> Path:
    doc = Document()
    configure_document(doc)
    add_cover(doc, images)

    heading(doc, "阅读说明", 1)
    paragraph(doc, "本手册说明坤伴在日常研发中的标准使用方法：从选择项目、组织上下文和发起会话，到处理任务、查看代码变化、创建变更集、审批交付和查看通知。界面会随版本持续迭代，实际可用入口以当前账号权限与管理员配置为准。")
    add_table(doc, ["角色", "主要工作", "常用入口"], [
        ["开发人员", "选择项目、组织上下文、与 AI 协作、查看运行和代码变化。", "聊天、任务面板、工作画布"],
        ["项目负责人", "关联任务、确认交付范围、组织会话协作与查看运行记录。", "任务面板、工作画布、代码交付"],
        ["代码审批人", "审核变更集、复核验证结果并执行受控交付。", "代码交付、项目 Git 状态"],
        ["系统管理员", "维护用户、项目授权、模型、第三方代理、Git 账号、推送和用量。", "工具与设置"],
    ], [1.15, 3.65, 2.3])
    heading(doc, "目录", 1)
    for item in ["1 平台概览", "2 快速开始", "3 聊天与代码协作", "4 任务面板与工作画布", "5 项目和代码库管理", "6 代码交付与仓库养护", "7 管理员设置", "8 安全边界与常见问题", "9 架构与产品演进参考", "10 知识库与多 Agent 协同建设"]:
        bullet(doc, item)

    heading(doc, "1 平台概览", 1)
    paragraph(doc, "坤伴是一套面向研发团队的 AI 协作工作台。它把项目代码、知识库、任务、Git、AI 会话和交付流程组织在同一个平台中，使团队能够在明确的项目与权限范围内完成分析、修改、验证、审批与推送。")
    add_figure(doc, images[2], "图 2  坤伴业务功能全景")
    add_table(doc, ["主链路", "用途", "关键控制"], [
        ["任务到会话", "将 Gitee 工作项、Issue 或本地任务带入对应项目会话。", "任务与项目授权、附件来源校验"],
        ["会话到改码", "让 Agent 读取已授权的代码、知识、Markdown 和附件并生成结果。", "受控上下文、执行模式、工具过程可见"],
        ["改码到交付", "将修改固化为变更集并完成审批、Git 推送与任务回写。", "校验、提交权限、SHA-256 再校验"],
        ["协作与通知", "在画布组织会话流程，并向钉钉项目群发送通知。", "项目绑定、运行记录、推送审计"],
    ], [1.4, 3.45, 2.25])
    heading(doc, "1.1 常见使用场景", 2)
    paragraph(doc, "坤伴以服务器上已挂载的代码库、已授权项目资料和 Git 记录作为协作事实基础。它不仅回答问题，也把修改、验证和交付的证据连续保留，便于研发与交付共同复核。")
    add_table(doc, ["场景", "可以这样提出需求", "产出与控制"], [
        ["开发新功能", "说明业务目标、影响模块、限制和验收条件，让 AI 先梳理方案与可能修改的文件。", "生成规划或实施说明后，在授权工作区修改；通过验证与变更集进入交付。"],
        ["规划文件", "让 AI 根据当前项目结构、已有 Markdown 与代码事实，整理接口方案、改造计划或评审材料。", "形成可引用的项目文档；不需要修改代码时可使用只读分析。"],
        ["排查 Bug 与历史差异", "询问某个 Bug 为什么出现，提供异常现象、时间或日志；也可要求比较某个日期或某次提交前的代码差异。", "返回关联文件、代码依据、Git 差异与待验证假设，避免脱离仓库事实猜测。"],
        ["简单开发和交付", "交付人员可直接描述范围明确的小改动，并说明验收方式。", "AI 生成修改和变更集；获授权人员审批、提交并推送，Git 保留提交与交付证据链供开发复查。"],
        ["复杂协同开发", "交付与开发共同上传最新需求变更文件或图片，并明确以服务器当前代码为准继续开发。", "用会话或工作画布分工、持续更新需求上下文，由开发人员复核差异、测试和最终交付。"],
    ], [1.25, 3.3, 2.55])

    heading(doc, "2 快速开始", 1)
    heading(doc, "2.1 登录并确认权限", 2)
    paragraph(doc, "使用管理员分配的账号登录。登录后先确认自己能看到正确的项目；如果无法登录或没有目标项目，请联系管理员开通或调整权限。不要在聊天、需求文件或截图中传递密码、令牌和连接串。")
    add_figure(doc, login_image, "图 3  坤伴平台登录入口", width=6.35)
    heading(doc, "2.2 选择项目并发起第一轮协作", 2)
    paragraph(doc, "首次使用时，请先确认自己已经被分配了至少一个项目。项目决定本次会话可使用的代码仓库、知识库、项目文档、运行配置和交付范围。")
    add_figure(doc, images[4], "图 4  聊天工作台入口")
    for text in [
        "在左侧选择聊天，确认当前会话所属项目。未选择项目时，AI 不会自动获得任何项目代码范围。",
        "在输入器中选择 Agent、模型和执行权限。日常分析可选只读；确需修改受信任项目时再选择工作区写入或完全控制。",
        "输入任务目标、预期结果和验收条件。涉及代码时，建议先让 AI 说明将读取哪些模块、准备修改哪些文件。",
        "回答完成后查看过程、差异、验证结果和 Git 状态；需要交付时进入变更集流程，而不是直接忽略审批。",
    ]:
        bullet(doc, text)

    heading(doc, "3 聊天与代码协作", 1)
    heading(doc, "3.1 选择项目和代码上下文", 2)
    paragraph(doc, "项目菜单用于选择当前项目。系统只展示当前用户有权使用的项目及其登记代码库。问题涉及指定模块时，请在提问中写明业务目标、入口路径或异常现象，避免让 AI 在过宽的范围内猜测。")
    add_figure(doc, images[5], "图 5  在聊天工作台中选择项目")
    heading(doc, "3.2 选择 Agent 模型和执行模式", 2)
    paragraph(doc, "聊天支持内置 Agent、Codex 本地代理和已配置的 DeepSeek Harness。不同代理支持的流式、图片、工具、代码写入能力不同；无法支持的能力应以界面状态为准。")
    add_figure(doc, images[6], "图 6  选择模型与 Agent")
    add_table(doc, ["执行模式", "适用场景", "使用提醒"], [
        ["只读分析", "理解代码、排查问题、设计方案、审查差异。", "不能修改工作区，适合先确认范围。"],
        ["工作区写入", "在指定项目工作区内修改已授权文件。", "发送前确认项目、代码库和文件范围。"],
        ["完全控制", "需要由 Codex 在受信任项目中直接执行操作。", "以服务端账户身份执行，不适用于未知或无关目录。"],
    ], [1.25, 3.0, 2.85])
    heading(doc, "3.3 提问与查看运行结果", 2)
    paragraph(doc, "建议用“目标、现状、限制、验收”描述问题。例如：说明需要修改的业务规则、不得影响的模块、希望执行的验证方式。过程区域用于查看模型阶段、工具调用和失败原因；长任务不要仅以页面等待状态判断是否失败。")
    add_figure(doc, images[7], "图 7  会话回答与过程信息")
    heading(doc, "3.4 使用项目 Markdown 和附件", 2)
    paragraph(doc, "可在聊天中引用项目 Markdown 文档，上传图片、文档或文本文件作为补充上下文。附件由后端保存为受控对象；不要在消息中粘贴令牌、连接串或生产数据。")
    add_figure(doc, images[10], "图 8  项目 Markdown 文档引用")
    add_figure(doc, images[11], "图 9  图片和文件附件入口")
    heading(doc, "3.5 查看差异 拉取和提交", 2)
    paragraph(doc, "聊天工具栏可查看项目与仓库的 Git 状态、待提交文件、差异和远端信息。重置更新会丢弃本地未提交修改，应先确认这些修改不需要保留。提交推送应在变更集审批完成后进行。")
    add_figure(doc, images[8], "图 10  聊天内的代码状态与 Git 操作")

    heading(doc, "4 任务面板与工作画布", 1)
    heading(doc, "4.1 从工作项创建会话", 2)
    paragraph(doc, "任务面板可新建本地任务、导入 CSV，或关联 Gitee 企业工作项和仓库 Issue。选择关联的 AiAgent 项目后，可创建带任务正文和合规图片附件的预填会话。批量创建时默认只生成草稿，不会自动执行。")
    heading(doc, "4.2 在工作画布组织会话", 2)
    paragraph(doc, "工作画布用于将已有会话作为节点组织起来。节点可以保存会话、预设模型、职责和 Skill；连线既可以投递上游结果，也可以表达执行依赖。双击节点打开右侧会话，布局和最近阅读位置会随画布保存。")
    add_figure(doc, images[13], "图 11  工作画布中的会话节点与检查器")
    bullet(doc, "投递关系：将上游节点的最新产出作为下游节点的受控输入。")
    bullet(doc, "流程关系：在上游条件完成后，推进可执行的下游节点，并保留运行记录。")
    bullet(doc, "画布不会复制或扩大原会话权限；仅能加入当前用户有权访问的会话。")

    heading(doc, "5 项目和代码库管理", 1)
    paragraph(doc, "项目是代码、知识和交付权限的组织单位。一个项目可登记多个仓库、运行配置和可在聊天中修改的配置文件。建议按实际业务系统创建项目，并避免将无关仓库混入同一项目。")
    add_figure(doc, images[18], "图 12  项目与代码库配置")
    heading(doc, "5.1 新建项目和挂载目录", 2)
    bullet(doc, "在项目与代码库中创建项目，并填写清晰的名称和说明。")
    bullet(doc, "挂载已有目录时，目录必须处于管理员配置的允许根目录内。")
    bullet(doc, "为代码库登记可运行配置和允许聊天编辑的文件，避免授予整个服务器目录。")
    heading(doc, "5.2 克隆远程仓库", 2)
    paragraph(doc, "选择目标项目后填写 HTTPS 仓库地址、分支和本地目录名，并选用已配置的 Git 账号。克隆过程通过实时通道显示；失败时检查 Git 账号权限、网络、仓库 URL 与目标目录是否已存在。")
    add_figure(doc, images[21], "图 13  克隆远程代码库")
    heading(doc, "5.3 Git 账号", 2)
    paragraph(doc, "Git 账号用于克隆、拉取和推送。令牌由服务端加密保存，界面不应展示或记录真实令牌。账户变更后建议先执行连接测试，再用于项目仓库。")
    add_figure(doc, images[22], "图 14  Git 账号配置")

    heading(doc, "6 代码交付与仓库养护", 1)
    heading(doc, "6.1 变更集交付", 2)
    paragraph(doc, "正式交付的标准流程是：创建变更集、执行固定 Git 校验、由具备代码提交权限的人员审批、再次复核工作区 SHA-256 指纹、提交并推送。审批后文件发生变化时，必须重新创建和审批变更集。")
    add_table(doc, ["步骤", "操作者", "结果"], [
        ["创建和校验", "开发人员或项目负责人", "形成修改摘要、验证状态和待审批变更集。"],
        ["审批", "具备代码提交权限的人员", "确认范围、风险、验证结果和提交信息。"],
        ["交付", "获授权审批人", "系统复核指纹后执行 Git 提交与推送，并记录交付结果。"],
    ], [1.2, 2.25, 3.65])
    heading(doc, "6.2 仓库养护", 2)
    paragraph(doc, "代码库养护用于对项目进行结构分析、优化建议、独立副本验证和维护分支交付。养护不绕过变更集审批：即使自动任务完成了分析或编译验证，仍需由具备权限的人员逐次批准后再推送。")
    add_figure(doc, images[15], "图 15  代码库养护设置与运行记录")

    heading(doc, "7 管理员设置", 1)
    paragraph(doc, "管理员负责把“人、项目、模型、代码库、外部平台”配置成可控的运行环境。普通用户只应看到自己被授权的项目和数据。")
    heading(doc, "7.1 模型与第三方代理", 2)
    paragraph(doc, "模型配置包含 LLM、Embedding 和相关供应商参数。第三方代理页面用于管理 Codex 模型策略和可用 Profile；请根据运行环境实际能力启用模型，不要把未验证的 CLI 当作可写入代理使用。")
    add_figure(doc, images[17], "图 16  模型配置入口")
    add_figure(doc, images[28], "图 17  第三方代理与模型策略")
    heading(doc, "7.2 用户权限和流量", 2)
    paragraph(doc, "用户管理可新增用户、分配项目、启停账户、重置密码和配置代码提交权限。历史会话与流量统计用于审计和容量判断；管理员不应通过界面或日志导出敏感附件与凭据。")
    add_figure(doc, images[24], "图 18  用户和项目授权管理")
    add_figure(doc, images[27], "图 19  用量统计")
    heading(doc, "7.3 钉钉推送", 2)
    paragraph(doc, "推送渠道与项目绑定后，可在 Git 推送成功时向对应项目群发送通知。群机器人接收任务请求时，会在后台定位项目并执行会话，最终将结果回传到原群。测试渠道前请确认群、机器人权限和服务器网络均可访问。")
    add_figure(doc, images[29], "图 20  钉钉推送渠道和项目绑定")

    heading(doc, "8 安全边界与常见问题", 1)
    add_table(doc, ["场景", "正确处理方式"], [
        ["需要 AI 修改代码", "先选择正确项目和仓库，限定问题范围；通过变更集校验与审批交付。"],
        ["需要完全控制", "只对受信任项目使用；理解它会按后端服务账户权限执行操作。"],
        ["看到 HTTP 500 或超时", "查看实时过程和后端日志，确认 Git/CLI/模型是否已完成，再决定是否重试。"],
        ["重置更新仓库", "先确认本地未提交文件无需保留；该操作会放弃本地修改并拉取远端。"],
        ["上传附件或引用外部任务", "不要上传密钥或生产敏感数据；外部内容应被视为待验证资料。"],
        ["无法访问项目或会话", "检查用户是否被分配项目、项目是否启用，以及当前登录会话是否有效。"],
    ], [1.85, 5.25])

    heading(doc, "9 架构与产品演进参考", 1)
    paragraph(doc, "坤伴 1.0 的重点是稳定交付闭环，后续版本优先提升可靠性、可观测性、上下文命中质量和交付效率。2.0 将建设知识资产中心、kun 资源接口、外部 Agent Adapter 与多 Agent 协作运行时。")
    add_figure(doc, images[3], "图 21  坤伴 1.0 到 2.0 迭代路径")
    paragraph(doc, "详细架构与规划请见本目录的资料索引。操作手册中的界面截图用于说明功能位置，实际界面、模型名单与权限选项会按部署配置有所不同。")

    heading(doc, "10 知识库与多 Agent 协同建设", 1)
    paragraph(doc, "2.0 的目标不是再增加一个孤立的知识库，而是把受控项目中的代码、任务、需求文件、会话结论和交付记录组织为可检索、可引用、可审计的知识资产。CPS Agent 与 APS Agent 应在同一项目权限边界内复用这些事实，从而提高自身分析、规划和协同能力。")
    heading(doc, "10.1 参考模式和业务流程", 2)
    paragraph(doc, "可参考 OpenViking 将资源、记忆和技能统一为可浏览上下文的思路：对资料保留摘要、概览和明细层级，按需加载，并记录检索轨迹。坤伴建议采用 kun URI 作为内部资源标识，例如 kun://project/{project}/code/{repo}/{path}@{commit}、kun://project/{project}/knowledge/{id}、kun://user/{user}/skill/{skill}。这是一项坤伴规划，不表示当前已经接入 OpenViking 或开放全部接口。")
    add_figure(doc, knowledge_diagrams["knowledge_flow"], "图 22  知识库与 CPS APS Agent 协同业务流程")
    add_table(doc, ["业务问题", "知识取用", "Agent 协同与可交付结果"], [
        ["CPS 出现 Bug 或需要比较某日期前的差异", "按项目、仓库、分支或提交定位代码、配置、日志和历史变更；返回来源与版本。", "CPS Agent 给出证据化诊断、影响范围和变更建议；需要改码时进入变更集与审批。"],
        ["APS 计划规则变化或需要评估业务影响", "检索计划规则、任务附件、需求版本、接口约束和历史结论。", "APS Agent 形成约束说明和方案；涉及 CPS 接口时投递已引用的上下文给 CPS Agent 复核。"],
        ["交付与开发共同处理复杂需求", "把新的需求文档或图片作为待验证资料挂到项目知识空间，并关联服务器代码事实。", "画布拆分分析、设计、修改和验证；最终保留引用、测试、审批和 Git 提交证据。"],
    ], [1.5, 2.1, 3.5])
    heading(doc, "10.2 建设阶段与验收重点", 2)
    paragraph(doc, "建设应从可控数据接入开始，而不是先追求自动写入。每个阶段都需要明确可验收指标，例如覆盖率、检索引用正确率、过期知识命中率、人工采纳率、跨 Agent 协同成功率，以及可追溯交付比例。")
    add_figure(doc, knowledge_diagrams["knowledge_roadmap"], "图 23  坤伴 2.0 知识库建设路径")
    add_table(doc, ["阶段", "优先交付", "通过标准"], [
        ["P0 数据接入", "项目、仓库、需求、附件和 Git 记录的增量采集；记录来源、版本、权限和有效期。", "可说明任一知识条目的来源和授权范围；失效或权限变更可以撤回访问。"],
        ["P1 知识可用", "知识分层、kun URI、检索引用、人工纠错和质量评估。", "回答能够给出关联资源和版本；不确定结论明确标记为待验证。"],
        ["P2 Agent 协同", "CPS 与 APS Adapter、共享上下文投递、会话记忆和跨系统影响复核。", "跨 Agent 任务不会越权；交接时保留输入、引用、结论和责任人。"],
        ["P3 自我演进", "高价值对话提炼、技能候选、灰度发布、回滚和质量指标闭环。", "知识或技能进入生产范围前经过评审、评估和可回滚验证。"],
    ], [1.05, 3.35, 2.7])
    paragraph(doc, "参考资料：OpenViking 官方文档 https://docs.openviking.ai/ 。其资源、记忆、技能统一组织和分层按需检索是本节的参考方向；坤伴将根据现有项目权限、Git 交付流程和部署约束实施自己的适配层。")

    path = output / "坤伴平台操作手册 2026-09-10.docx"
    doc.core_properties.title = "坤伴平台操作手册"
    doc.core_properties.subject = "坤伴平台操作与管理指南"
    doc.core_properties.author = "坤伴平台"
    doc.save(path)
    return path


def html_escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def figure_html(number: int, title: str) -> str:
    return f'''<figure class="shot"><img src="assets/source-screenshots/screen-{number:02d}.png" alt="{html_escape(title)}" loading="lazy"/><figcaption>{html_escape(title)}</figcaption></figure>'''


def figure_html_file(relative_path: str, title: str) -> str:
    return f'''<figure class="shot"><img src="{html_escape(relative_path)}" alt="{html_escape(title)}" loading="lazy"/><figcaption>{html_escape(title)}</figcaption></figure>'''


def build_html(output: Path) -> Path:
    figures = {
        "architecture": figure_html(1, "坤伴平台当前系统架构"),
        "business": figure_html(2, "坤伴业务功能全景"),
        "roadmap": figure_html(3, "坤伴 1.0 到 2.0 迭代路径"),
        "login": figure_html_file("assets/login-screen.png", "坤伴平台登录入口"),
        "knowledge_flow": figure_html_file("assets/knowledge-cps-aps-collaboration.png", "知识库与 CPS APS Agent 协同业务流程"),
        "knowledge_roadmap": figure_html_file("assets/knowledge-2.0-roadmap.png", "坤伴 2.0 知识库建设路径"),
        "chat": figure_html(4, "聊天工作台入口"),
        "project": figure_html(5, "在聊天中选择项目"),
        "agent": figure_html(6, "选择模型和 Agent"),
        "answer": figure_html(7, "会话回答和过程信息"),
        "git": figure_html(8, "聊天内 Git 状态和代码操作"),
        "markdown": figure_html(10, "项目 Markdown 文档引用"),
        "attachment": figure_html(11, "图片和文件附件入口"),
        "canvas": figure_html(13, "工作画布中的会话节点"),
        "repository": figure_html(18, "项目与代码库配置"),
        "clone": figure_html(21, "克隆远程代码库"),
        "gitaccount": figure_html(22, "Git 账号配置"),
        "maintenance": figure_html(15, "代码库养护设置与运行记录"),
        "models": figure_html(17, "模型配置入口"),
        "agentprovider": figure_html(28, "第三方代理和模型策略"),
        "admin": figure_html(24, "用户和项目授权管理"),
        "usage": figure_html(27, "用量统计"),
        "push": figure_html(29, "钉钉推送渠道和项目绑定"),
    }
    sections = f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"/><meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>坤伴平台操作手册</title><style>
:root{{--bg:#f6f8fc;--paper:#fff;--ink:#172033;--muted:#60708a;--line:#dbe5f2;--blue:#2563eb;--navy:#173b69;--soft:#eef5ff}}*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:var(--bg);color:var(--ink);font:16px/1.7 "Microsoft YaHei","Noto Sans CJK SC",Arial,sans-serif}}.hero{{background:linear-gradient(120deg,#123867,#2563eb 55%,#6d28d9);color:#fff;padding:72px max(30px,calc((100vw - 1240px)/2));}}.hero h1{{margin:0;font-size:clamp(30px,4vw,52px);letter-spacing:.02em}}.hero p{{max-width:760px;margin:14px 0 0;opacity:.9}}.shell{{max-width:1240px;margin:0 auto;display:grid;grid-template-columns:250px minmax(0,1fr);gap:32px;padding:32px 24px 72px}}nav{{position:sticky;top:18px;align-self:start;background:var(--paper);border:1px solid var(--line);border-radius:16px;padding:14px;box-shadow:0 8px 24px rgba(15,35,60,.06)}}nav strong{{display:block;padding:4px 10px 10px}}nav a{{display:block;padding:7px 10px;color:#314764;text-decoration:none;border-radius:8px;font-size:14px}}nav a:hover{{background:var(--soft);color:var(--blue)}}main{{min-width:0}}.intro,.section{{background:var(--paper);border:1px solid var(--line);border-radius:18px;padding:30px 34px;margin-bottom:22px;box-shadow:0 8px 24px rgba(15,35,60,.045)}}h2{{font-size:25px;line-height:1.3;margin:0 0 16px;color:#111827}}h3{{font-size:19px;margin:28px 0 10px;color:#111827}}p{{margin:10px 0}}.eyebrow{{font-size:12px;letter-spacing:.14em;color:var(--blue);font-weight:700}}.grid{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px;margin:18px 0}}.card{{border:1px solid var(--line);border-radius:12px;padding:15px;background:#fff}}.card b{{display:block;margin-bottom:5px}}.flow{{display:flex;flex-wrap:wrap;gap:8px;margin:18px 0}}.flow span{{padding:7px 11px;background:var(--soft);color:#18468a;border-radius:8px;font-size:14px;font-weight:700}}.flow span:not(:last-child)::after{{content:'→';margin-left:13px;color:#8094b4}}table{{width:100%;border-collapse:collapse;margin:16px 0;overflow:hidden;border-radius:10px}}th,td{{border:1px solid var(--line);padding:10px 12px;text-align:left;vertical-align:top;font-size:14px}}th{{background:var(--navy);color:#fff}}tr:nth-child(even) td{{background:#f8fbff}}.shot{{margin:22px 0;border:1px solid var(--line);border-radius:12px;overflow:hidden;background:#fff}}.shot img{{display:block;width:100%;height:auto;cursor:zoom-in}}figcaption{{padding:9px 14px;color:var(--muted);font-size:13px;background:#fafcff}}details{{border:1px solid var(--line);border-radius:10px;padding:11px 14px;margin:10px 0}}summary{{cursor:pointer;font-weight:700}}.checklist label{{display:block;padding:8px 0;border-bottom:1px dashed var(--line)}}.checklist input{{margin-right:8px;accent-color:var(--blue)}}.toolrow{{display:flex;gap:10px;align-items:center;margin-top:20px}}input[type=search]{{width:100%;padding:10px 12px;border:1px solid var(--line);border-radius:9px;font:inherit}}button{{border:1px solid var(--line);background:#fff;border-radius:8px;padding:8px 11px;cursor:pointer;color:var(--ink)}}button:hover{{background:var(--soft)}}.hidden{{display:none!important}}#modal{{position:fixed;inset:0;background:rgba(8,15,30,.78);display:none;align-items:center;justify-content:center;padding:20px;z-index:10}}#modal.open{{display:flex}}#modal img{{max-width:96vw;max-height:92vh;border-radius:8px;background:#fff}}@media(max-width:850px){{.hero{{padding:44px 22px}}.shell{{display:block;padding:18px 14px 44px}}nav{{position:relative;top:auto;margin-bottom:18px;display:flex;gap:2px;overflow:auto;white-space:nowrap}}nav strong{{display:none}}nav a{{display:inline-block;padding:8px}}.intro,.section{{padding:22px 18px}}.grid{{grid-template-columns:1fr}}table{{display:block;overflow:auto}}}}@media print{{nav,.toolrow{{display:none}}.shell{{display:block;padding:0}}.intro,.section{{box-shadow:none;border:0;margin:0 0 16px;padding:0}}.hero{{padding:30px;background:#fff;color:#000}}}}
</style></head><body><header class="hero"><div class="eyebrow" style="color:#dbeafe">KUNBUDDY PLATFORM</div><h1>坤伴平台操作手册</h1><p>面向研发协作、代码交付与平台管理。按项目范围组织 AI 会话、任务、知识、代码和交付过程。</p></header>
<div class="shell"><nav><strong>操作目录</strong><a href="#overview">平台概览</a><a href="#quickstart">快速开始</a><a href="#chat">聊天与代码协作</a><a href="#tasks">任务与工作画布</a><a href="#repository">项目与代码库</a><a href="#delivery">交付与养护</a><a href="#settings">管理员设置</a><a href="#safety">安全与常见问题</a><a href="#roadmap">架构与演进</a><a href="#knowledge">知识库与多 Agent</a></nav><main>
<section class="intro"><div class="eyebrow">使用前先确认</div><h2>项目范围和权限决定可用能力</h2><p>选择项目后，系统才会提供对应代码库、知识库、项目文档、运行配置和交付范围。每一次改码和交付都应在授权项目内完成。</p><div class="flow"><span>任务或问题</span><span>AI 协作</span><span>受控改码</span><span>审批交付</span><span>Git 与通知</span></div><div class="checklist" id="checklist"><label><input type="checkbox" data-key="project"/>我已确认当前项目和代码库正确</label><label><input type="checkbox" data-key="scope"/>我已说明任务目标、限制和验收方式</label><label><input type="checkbox" data-key="delivery"/>我知道正式提交需经过变更集审批</label></div></section>
<section class="section searchable" id="overview"><div class="eyebrow">01</div><h2>平台概览</h2><p>坤伴是一套面向研发团队的 AI 协作工作台，将项目代码、知识库、任务、Git、AI 会话和交付流程组织在同一平台中。</p>{figures['architecture']}{figures['business']}<div class="grid"><div class="card"><b>开发人员</b>选择项目和上下文，与 AI 协作，查看运行、差异和验证结果。</div><div class="card"><b>项目负责人</b>导入工作项、组织会话和画布、确认交付范围。</div><div class="card"><b>代码审批人</b>审核变更集、复核指纹与验证结果，执行受控交付。</div><div class="card"><b>系统管理员</b>维护用户、项目授权、模型、Git 账号、推送和用量。</div></div><h3>坤伴可以帮助完成什么</h3><p>坤伴以服务器上已挂载的代码库、已授权项目资料和 Git 记录作为事实基础。它把查询、改动、验证和交付记录连续保留，便于研发与交付共同复核。</p><table><thead><tr><th>场景</th><th>示例提问或操作</th><th>产出与控制</th></tr></thead><tbody><tr><td>开发新功能</td><td>说明业务目标、影响模块、限制和验收条件。</td><td>先形成方案，再在授权工作区修改并验证。</td></tr><tr><td>规划文件</td><td>根据当前项目结构、已有 Markdown 和代码事实，整理接口方案或改造计划。</td><td>形成可引用项目文档；无需改码时使用只读分析。</td></tr><tr><td>Bug 与历史差异</td><td>询问某个 Bug 的原因，或比较某个日期、某次提交前的代码差异。</td><td>返回关联文件、代码依据、Git 差异和待验证假设。</td></tr><tr><td>简单开发和交付</td><td>交付人员直接描述范围明确的小改动和验收方式。</td><td>变更集经授权审批后提交推送，Git 保留证据链供开发复查。</td></tr><tr><td>复杂协同开发</td><td>交付与开发共同上传最新需求变更文件或图片，以服务器当前代码为准继续开发。</td><td>在会话或工作画布分工，持续更新上下文并复核测试和差异。</td></tr></tbody></table></section>
<section class="section searchable" id="quickstart"><div class="eyebrow">02</div><h2>快速开始</h2><h3>登录并确认权限</h3><p>使用管理员分配的账号登录。登录后先确认自己能看到正确的项目；无法登录或没有目标项目时，请联系管理员开通或调整权限。不要在聊天、需求文件或截图中传递密码、令牌和连接串。</p>{figures['login']}<h3>选择项目并发起第一轮协作</h3>{figures['chat']}<ol><li>登录后确认自己已被分配到目标项目。</li><li>在聊天中选择项目、Agent、模型和执行权限。</li><li>描述目标、现状、限制与验收方式；代码问题尽量给出模块或入口。</li><li>阅读过程、差异和验证结果；需要提交时进入变更集流程。</li></ol></section>
<section class="section searchable" id="chat"><div class="eyebrow">03</div><h2>聊天与代码协作</h2><h3>选择项目和代码上下文</h3><p>系统只展示当前用户有权访问的项目与仓库。项目不是普通筛选器，而是本次会话的可用资源范围。</p>{figures['project']}<h3>选择 Agent 模型和执行模式</h3><p>内置 Agent、Codex 本地代理和已配置的 DeepSeek Harness 的能力不同。日常分析优先使用只读；需要修改受信任项目时再选择工作区写入或完全控制。</p>{figures['agent']}<table><thead><tr><th>执行模式</th><th>适用场景</th><th>提醒</th></tr></thead><tbody><tr><td>只读分析</td><td>代码理解、问题排查、方案设计、差异审查</td><td>不能修改工作区</td></tr><tr><td>工作区写入</td><td>修改当前项目已授权文件</td><td>发送前确认范围</td></tr><tr><td>完全控制</td><td>由 Codex 在受信任项目执行操作</td><td>以服务端账户权限执行</td></tr></tbody></table><h3>提问、过程和附件</h3><p>建议使用目标、现状、限制和验收描述任务。过程区会显示运行阶段、工具调用和失败信息；长任务不要只按页面等待时间判断结果。</p>{figures['answer']}{figures['markdown']}{figures['attachment']}<h3>查看差异 拉取和提交</h3><p>可查看待提交文件、差异和远端状态。重置更新会丢弃本地未提交修改，提交推送应在变更集审批完成后执行。</p>{figures['git']}</section>
<section class="section searchable" id="tasks"><div class="eyebrow">04</div><h2>任务面板与工作画布</h2><h3>从工作项创建会话</h3><p>任务面板支持本地任务、CSV 导入和 Gitee 工作项或 Issue。关联目标项目后，可以生成预填会话；批量创建时默认只生成草稿，不会自动执行。</p><h3>在工作画布组织会话</h3><p>画布将已有会话作为节点。节点可保存模型预设、职责和 Skill；连线可投递上游结果或表达执行依赖。双击节点可打开右侧会话。</p>{figures['canvas']}<details><summary>画布连线如何选择</summary><p><b>投递关系</b>适用于把上游产出作为下游受控输入。<br/><b>流程关系</b>适用于建立前置条件和运行推进关系。两种关系都不会绕过原会话的项目权限和运行配额。</p></details></section>
<section class="section searchable" id="repository"><div class="eyebrow">05</div><h2>项目与代码库管理</h2><p>一个项目可登记多个仓库、运行配置和可由聊天修改的配置文件。请按实际业务系统划分项目，不要混入无关仓库。</p>{figures['repository']}<h3>克隆远程仓库</h3><p>选择项目后填写 HTTPS 地址、分支和本地目录名，并选用已配置的 Git 账号。失败时检查仓库地址、账号权限、网络与目标目录。</p>{figures['clone']}<h3>Git 账号</h3><p>Git 账号用于克隆、拉取和推送。令牌由服务端加密保存，不应在聊天、文档或日志中记录真实令牌。</p>{figures['gitaccount']}</section>
<section class="section searchable" id="delivery"><div class="eyebrow">06</div><h2>代码交付与仓库养护</h2><p>正式交付遵循创建变更集、固定 Git 校验、人工审批、SHA-256 再校验和提交推送的顺序。审批后的工作区如果变化，必须重新创建和审批变更集。</p><table><thead><tr><th>步骤</th><th>操作者</th><th>结果</th></tr></thead><tbody><tr><td>创建和校验</td><td>开发人员或项目负责人</td><td>形成变更摘要、验证状态和待审批记录</td></tr><tr><td>审批</td><td>具备提交权限的人员</td><td>确认范围、风险、验证和提交信息</td></tr><tr><td>交付</td><td>获授权审批人</td><td>再次复核指纹后提交并推送</td></tr></tbody></table><h3>代码库养护</h3><p>养护可在独立副本内分析、优化和验证，但不绕过变更集审批；维护分支推送仍需逐次审批。</p>{figures['maintenance']}</section>
<section class="section searchable" id="settings"><div class="eyebrow">07</div><h2>管理员设置</h2><h3>模型和第三方代理</h3><p>管理员维护模型供应商、模型目录和第三方代理策略。仅启用已经过运行环境验证的模型和 CLI 能力。</p>{figures['models']}{figures['agentprovider']}<h3>用户权限和用量</h3><p>用户管理用于分配项目、启停账户、重置密码和配置代码提交权限。流量统计用于审计与容量判断。</p>{figures['admin']}{figures['usage']}<h3>钉钉推送</h3><p>项目与推送渠道绑定后，Git 推送成功可通知项目群。群机器人收到请求后，将定位项目并在后台执行会话，再回传结果。</p>{figures['push']}</section>
<section class="section searchable" id="safety"><div class="eyebrow">08</div><h2>安全边界与常见问题</h2><details open><summary>需要 AI 修改代码</summary><p>先选择正确项目和仓库，限定问题范围；正式提交走变更集校验与审批。</p></details><details><summary>出现 HTTP 500 或超时</summary><p>先看实时过程、运行状态和后端日志，确认 Git、CLI 或模型是否已经执行完成，再决定是否重试。</p></details><details><summary>重置更新仓库</summary><p>该操作会放弃本地未提交修改并拉取远端代码。执行前确认修改已经不需要保留，或已通过变更集交付。</p></details><details><summary>上传附件或引用外部任务</summary><p>不要上传令牌、连接串或生产敏感数据。外部 Issue、图片和 OCR 文本应作为待验证资料，不应改变工具权限。</p></details></section>
<section class="section searchable" id="roadmap"><div class="eyebrow">09</div><h2>架构与产品演进</h2><p>1.0 的目标是稳定的研发交付闭环，1.x 优先提升可靠性、可观测性、上下文质量和交付效率。2.0 将建设知识资产、kun 资源接口、外部 Agent Adapter 与多 Agent 协同运行时。</p>{figures['roadmap']}<p>详细说明参见本目录的资料索引和已整理的架构与路线图 Markdown 文档。</p></section>
<section class="section searchable" id="knowledge"><div class="eyebrow">10</div><h2>知识库与多 Agent 协同建设</h2><p>坤伴 2.0 的知识库不是孤立资料库，而是将已授权项目中的代码、任务、需求文件、会话结论和交付记录组织为可检索、可引用、可审计的知识资产。CPS Agent 与 APS Agent 在统一项目权限内使用这些事实，提高问题诊断、规划和协同能力。</p><h3>参考模式</h3><p>参考 OpenViking 的上下文组织思路，可将资源、记忆与技能作为统一、可浏览的上下文，并按摘要、概览和明细分层加载。坤伴建议使用 <code>kun://project/&#123;project&#125;/code/&#123;repo&#125;/&#123;path&#125;@&#123;commit&#125;</code>、<code>kun://project/&#123;project&#125;/knowledge/&#123;id&#125;</code> 和 <code>kun://user/&#123;user&#125;/skill/&#123;skill&#125;</code> 标识内部资源。以下是坤伴建设规划，并不表示当前已经接入 OpenViking 或开放全部接口。</p>{figures['knowledge_flow']}<h3>业务示例</h3><table><thead><tr><th>业务问题</th><th>知识取用</th><th>Agent 协同与可交付结果</th></tr></thead><tbody><tr><td>CPS 出现 Bug 或需比较历史差异</td><td>按项目、仓库、分支或提交定位代码、配置、日志和 Git 历史，返回来源与版本。</td><td>CPS Agent 给出证据化诊断、影响范围和变更建议；改码进入变更集与审批。</td></tr><tr><td>APS 计划规则变化或需评估影响</td><td>检索计划规则、任务附件、需求版本、接口约束和历史结论。</td><td>APS Agent 输出约束说明和方案；涉及 CPS 接口时将已引用上下文投递给 CPS Agent 复核。</td></tr><tr><td>交付与开发共同处理复杂需求</td><td>新需求文件或图片先作为待验证资料挂到项目知识空间，并关联服务器代码事实。</td><td>画布拆分分析、设计、修改和验证；保留引用、测试、审批和 Git 提交证据。</td></tr></tbody></table><h3>建设路径与验收重点</h3><p>先完成可控接入，再建设可用知识，然后开展 Agent 协同，最后才扩大自动化与自我演进。每阶段都需要度量覆盖率、检索引用正确率、过期知识命中率、人工采纳率、协同成功率和可追溯交付比例。</p>{figures['knowledge_roadmap']}<table><thead><tr><th>阶段</th><th>优先交付</th><th>通过标准</th></tr></thead><tbody><tr><td>P0 数据接入</td><td>项目、仓库、需求、附件和 Git 记录增量采集；记录来源、版本、权限、有效期。</td><td>每条知识能说明来源和授权范围；权限变化可撤回访问。</td></tr><tr><td>P1 知识可用</td><td>知识分层、kun URI、检索引用、人工纠错和质量评估。</td><td>回答能给出资源与版本；不确定结论标记待验证。</td></tr><tr><td>P2 Agent 协同</td><td>CPS 与 APS Adapter、共享上下文投递、会话记忆和跨系统影响复核。</td><td>交接不越权，且保留输入、引用、结论与责任人。</td></tr><tr><td>P3 自我演进</td><td>高价值对话提炼、技能候选、灰度发布、回滚和指标闭环。</td><td>知识或技能进入生产范围前完成评审、评估和回滚验证。</td></tr></tbody></table><p>参考：<a href="https://docs.openviking.ai/" target="_blank" rel="noreferrer">OpenViking 官方文档</a>。坤伴将结合当前项目权限、Git 交付流程和部署约束建设自身适配层。</p></section></main></div><div id="modal" aria-label="图片预览"><img alt="图片预览"/></div><script>
const modal=document.querySelector('#modal'), modalImg=modal.querySelector('img');document.querySelectorAll('.shot img').forEach(img=>img.addEventListener('click',()=>{{modalImg.src=img.src;modal.classList.add('open')}}));modal.addEventListener('click',()=>modal.classList.remove('open'));document.querySelectorAll('#checklist input').forEach(input=>{{input.checked=localStorage.getItem('kun-manual-'+input.dataset.key)==='1';input.addEventListener('change',()=>localStorage.setItem('kun-manual-'+input.dataset.key,input.checked?'1':'0'))}});const search=document.createElement('input');search.type='search';search.placeholder='搜索手册内容';document.querySelector('nav').append(search);search.addEventListener('input',e=>{{const q=e.target.value.trim().toLowerCase();document.querySelectorAll('.searchable').forEach(s=>s.classList.toggle('hidden',q&&!s.innerText.toLowerCase().includes(q)))}});
</script></body></html>'''
    path = output / "坤伴平台操作手册 2026-09-10.html"
    path.write_text(sections, encoding="utf-8")
    return path


def build_readme(output: Path) -> Path:
    content = """# 坤伴平台手册交付目录

本目录汇总坤伴平台的优化版操作手册、可交互 HTML 手册、原始手册截图，以及当前架构和产品迭代资料。

## 交付物

- `坤伴平台操作手册 2026-09-10.docx`：适合下载、打印和线下交付的优化版 Word 手册。
- `坤伴平台操作手册 2026-09-10.html`：适合浏览器阅读，支持目录跳转、内容搜索、图片放大和本地保存的使用前检查项。
- `assets/source-screenshots/`：从用户提供的原手册中保留的界面截图，用于说明操作位置。
- `assets/knowledge-cps-aps-collaboration.png` 与 `assets/knowledge-2.0-roadmap.png`：本次生成的知识库和多 Agent 协同建设图，用于说明坤伴 2.0 规划。

## 关联资料

- [当前架构与业务功能分析](参考资料/2026-09-10-kunbuddy-aiagent-architecture-and-capability-analysis.md)
- [1.0 到 2.0 迭代规划](参考资料/2026-09-10-kunbuddy-1-to-2-roadmap.md)
- [系统架构图](参考资料/images/2026-09-10-kunbuddy-system-architecture.svg)
- [业务功能全景图](参考资料/images/2026-09-10-kunbuddy-business-capabilities.svg)
- [2.0 规划图](参考资料/images/2026-09-10-kunbuddy-2.0-roadmap.svg)
- [OpenViking 官方文档](https://docs.openviking.ai/)：知识资源、记忆与技能的统一上下文组织参考。

## 维护建议

界面、模型名单、部署地址和用户权限会随环境变化。功能改动后，请同时更新操作手册对应章节、截图和关联架构资料；不要在手册、截图、HTML 或示例中写入真实令牌、密码、连接串或生产数据。
"""
    path = output / "README.md"
    path.write_text(content, encoding="utf-8")
    return path


def collect_reference_materials(output: Path) -> None:
    """Make the delivery folder self-contained with the architecture materials."""
    source = output.parent / "architecture"
    target = output / "参考资料"
    target.mkdir(parents=True, exist_ok=True)
    for filename in (
        "2026-09-10-kunbuddy-aiagent-architecture-and-capability-analysis.md",
        "2026-09-10-kunbuddy-1-to-2-roadmap.md",
    ):
        shutil.copy2(source / filename, target / filename)
    source_images = source / "images"
    target_images = target / "images"
    target_images.mkdir(parents=True, exist_ok=True)
    for filename in (
        "2026-09-10-kunbuddy-system-architecture.svg",
        "2026-09-10-kunbuddy-business-capabilities.svg",
        "2026-09-10-kunbuddy-2.0-roadmap.svg",
    ):
        shutil.copy2(source_images / filename, target_images / filename)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--login-image", type=Path, required=True)
    parser.add_argument("--inspect", action="store_true")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.inspect:
        inspect(args.source, args.output)
        return
    assets = args.output / "assets"
    image_paths = extract_source_images(args.source, assets)
    images = {index: path for index, path in enumerate(image_paths, start=1)}
    login_image = assets / "login-screen.png"
    shutil.copy2(args.login_image, login_image)
    knowledge_diagrams = build_knowledge_diagrams(assets)
    docx_path = build_docx(args.output, images, login_image, knowledge_diagrams)
    html_path = build_html(args.output)
    collect_reference_materials(args.output)
    readme_path = build_readme(args.output)
    print(json.dumps({"docx": str(docx_path), "html": str(html_path), "readme": str(readme_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
