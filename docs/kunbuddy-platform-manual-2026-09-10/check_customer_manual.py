"""Render PDF and check the customer brochure (playwright + pymupdf).

Usage: python check_customer_manual.py --browser PATH_TO_CHROME
Review images are written to the OS temporary directory, outside the public pack.
"""
import argparse
import re
import tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pymupdf as fitz
from docx import Document
from playwright.sync_api import sync_playwright

from build_customer_manual import OUT, NAME, SOURCE


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--browser", required=True)
    args = parser.parse_args()
    review = Path(tempfile.gettempdir()) / "kunbuddy-customer-manual-review"
    review.mkdir(exist_ok=True)
    headings = re.findall(r"^## (.+)$", SOURCE.read_text(encoding="utf-8"), re.M)
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=args.browser, headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 1080}, device_scale_factor=1)
        external = []
        page.on("request", lambda req: external.append(req.url) if req.url.startswith(("https:", "http:")) else None)
        page.goto((OUT / (NAME + ".html")).as_uri(), wait_until="networkidle")
        page.evaluate("document.fonts.ready")
        assert page.locator(".page").count() == 12
        assert page.locator(".toc a").count() == 10
        for i in range(1, 11):
            assert page.locator(f"#chapter-{i}").count() == 1
        page.locator(".toc a").last.click()
        page.wait_for_function("location.hash === '#chapter-10'")
        page.locator("#cover").screenshot(path=str(review / "cover.png"))
        page.pdf(path=str(OUT / (NAME + ".pdf")), print_background=True, prefer_css_page_size=True, display_header_footer=False)
        page.set_viewport_size({"width": 390, "height": 844})
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), "Mobile overflow"
        page.locator("#chapter-6").screenshot(path=str(review / "mobile-delivery.png"))
        assert not external, external
        browser.close()
    pdf = fitz.open(OUT / (NAME + ".pdf"))
    assert len(pdf) == 12, f"Unexpected PDF page count: {len(pdf)}"
    for i, title in enumerate(headings, 2):
        text = pdf[i].get_text()
        assert title[:2] in text
        assert title.split("｜")[0][3:] in text
        assert len(text) > 200, f"Sparse page {i+1}"
    for idx in (0, 2, 4, 7, 11):
        pdf[idx].get_pixmap(matrix=fitz.Matrix(1, 1)).save(review / f"pdf-page-{idx+1:02d}.png")
    word = Document(OUT / (NAME + ".docx"))
    assert len([p for p in word.paragraphs if p.style.name == "Heading 1"]) == 11
    assert len(word.tables) == 11
    assert not word.inline_shapes
    with ZipFile(OUT / (NAME + ".docx")) as z:
        rels = z.read("word/_rels/document.xml.rels").decode()
        assert 'TargetMode="External"' not in rels
        xml = z.read("word/document.xml").decode()
        assert xml.count("w:bookmarkStart") == 10
        assert xml.count("w:hyperlink") == 20
    public_text = SOURCE.read_text(encoding="utf-8") + (OUT / (NAME + ".html")).read_text(encoding="utf-8")
    for forbidden in ("D:\\", "C:\\Users", "df7e4d2", "yun_kun", "backed/", "localhost", "sk-", "Bearer ", "待填写", "TODO"):
        assert forbidden not in public_text, forbidden
    assert "\ufffd" not in public_text
    with ZipFile(OUT / "坤伴Agent平台客户介绍资料包.zip", "w", ZIP_DEFLATED) as archive:
        for ext in (".docx", ".pdf", ".html", ".md"):
            archive.write(OUT / (NAME + ext), NAME + ext)
    print(f"PASS: 12 PDF pages; 10 chapters; Word tables/bookmarks; mobile layout; offline assets; public-content checks. Review: {review}")
    print("Word package checked structurally; Word pagination requires a compatible office renderer.")


if __name__ == "__main__":
    main()
