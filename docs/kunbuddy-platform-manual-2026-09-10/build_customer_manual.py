"""Build the customer brochure from one UTF-8 Markdown source.

Requirements: python-docx. Optional PDF: print the HTML with a local browser.
No source screenshots, private project data, or external assets are embedded.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "对外介绍-2026-09-18"
NAME = "坤伴Agent协同研发平台介绍手册"
SOURCE = OUT / (NAME + ".md")
NAVY = "143049"
TEAL = "087F8C"


def blocks(text):
    lines = text.strip().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            yield "table", rows
            continue
        for prefix, kind in [("### ", "h3"), ("> ", "quote"), ("- ", "bullet")]:
            if line.startswith(prefix):
                yield kind, line[len(prefix):]
                break
        else:
            yield "flow" if " → " in line else "p", line
        i += 1


def set_font(style, size, color=NAVY, bold=False):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")


def shade(cell, color):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(el)


def bookmark(paragraph, name, number):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def internal_link(paragraph, title, anchor):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def build_docx(chapters):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.05), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.05), Cm(2.05)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)
    section.different_first_page_header_footer = True
    for name, size, color, bold in [
        ("Normal", 10.5, NAVY, False), ("Title", 32, NAVY, True),
        ("Subtitle", 16, TEAL, False), ("Heading 1", 21, NAVY, True),
        ("Heading 2", 12, TEAL, True), ("List Bullet", 10.5, NAVY, False),
        ("Caption", 9, "62788A", False),
    ]:
        set_font(doc.styles[name], size, color, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.25
    normal.space_after = Pt(8)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.space_after = Pt(10)
        doc.styles[name].paragraph_format.space_before = Pt(10)
    header = section.header.paragraphs[0]
    header.text = "KUNBUDDY  /  坤伴                        AGENT 协同研发平台 · 客户介绍"
    set_font(doc.styles["Header"], 8, "62788A")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("2026.09  ·  客户介绍版     /     ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_font(doc.styles["Footer"], 8, "62788A")
    doc.core_properties.title = "坤伴 Agent 协同研发平台介绍手册"
    doc.core_properties.subject = "客户介绍：协同效率、项目事实与可核验交付"
    doc.core_properties.author = "坤伴"
    doc.core_properties.keywords = "Agent, 协同研发, 客户介绍, 交付治理"
    doc.core_properties.comments = "客户介绍版；能力范围以项目约定与部署验收为准。"
    doc.add_paragraph("KUNBUDDY  /  坤伴", "Subtitle")
    doc.add_paragraph("客户介绍手册   /   2026.09", "Caption")
    doc.add_paragraph("").paragraph_format.space_after = Pt(45)
    doc.add_paragraph("让需求推进更连贯\n让交付依据更清晰", "Title")
    doc.add_paragraph("Agent 协同研发平台", "Subtitle")
    p = doc.add_paragraph("连接任务、项目事实、Agent 会话与代码交付，\n让每一阶段都有具体产出，每一次交付都有核对依据。")
    p.paragraph_format.space_before = Pt(22)
    doc.add_paragraph("").paragraph_format.space_after = Pt(30)
    table = doc.add_table(rows=1, cols=3)
    for cell, title in zip(table.rows[0].cells, ["协作更连贯\n阶段产出有承接", "过程可回看\n状态与依据可核对", "交付有控制\n人工审批与版本记录"]):
        shade(cell, "EAF5F5")
        cell.text = title
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph("面向业务负责人、项目负责人、研发与信息化团队", "Caption").paragraph_format.space_before = Pt(50)
    doc.add_paragraph("版本日期：2026-09-18。本手册依据当前代码实现整理；具体开通范围、模型服务及实施安排以项目约定和部署验收为准。", "Caption")
    doc.add_page_break()
    doc.add_heading("阅读目录", 1)
    doc.add_paragraph("从平台价值开始，了解协作方式、交付机制与试点路径。")
    for i, (title, _) in enumerate(chapters, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(13)
        internal_link(p, title.replace("｜", "  /  "), f"chapter_{i}")
    doc.add_heading("建议阅读路径", 2)
    doc.add_paragraph("业务与决策负责人：01、02、09、10\n项目与交付负责人：03、06、07、09\n研发与信息化团队：04、05、06、08")
    doc.add_paragraph("文中流程和场景用于说明平台的使用方式，不代表客户案例或量化效果承诺。", "Caption")
    for i, (title, body) in enumerate(chapters, 1):
        doc.add_page_break()
        h = doc.add_heading(title.replace("｜", "\n"), 1)
        bookmark(h, f"chapter_{i}", i)
        for kind, value in blocks(body):
            if kind == "table":
                table = doc.add_table(rows=1, cols=len(value[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Light Shading Accent 1"
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row_idx, values in enumerate(value):
                    cells = table.rows[0].cells if row_idx == 0 else table.add_row().cells
                    row = table.rows[row_idx]
                    no_split = OxmlElement("w:cantSplit")
                    row._tr.get_or_add_trPr().append(no_split)
                    for cell, text in zip(cells, values):
                        cell.text = text
                        shade(cell, NAVY if row_idx == 0 else ("EFF6F8" if row_idx % 2 else "FFFFFF"))
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        for p in cell.paragraphs:
                            p.paragraph_format.space_before = Pt(5)
                            p.paragraph_format.space_after = Pt(5)
                            p.paragraph_format.line_spacing = 1.15
                            for run in p.runs:
                                run.font.size = Pt(9)
                                run.font.bold = row_idx == 0
                                run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_idx == 0 else NAVY)
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            elif kind == "h3":
                doc.add_heading(value, 2)
            elif kind in ("quote", "flow"):
                p = doc.add_paragraph(value)
                p.paragraph_format.space_before = Pt(6)
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string(TEAL)
                    run.font.bold = True
            else:
                doc.add_paragraph(value, "List Bullet" if kind == "bullet" else None)
    doc.save(OUT / (NAME + ".docx"))


CSS = """
:root{--ink:#143049;--muted:#607889;--teal:#087f8c;--paper:#fff;--line:#d8e5eb}
*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#eaf0f4;color:var(--ink);font-family:'Microsoft YaHei','Noto Sans CJK SC',sans-serif;font-size:15px;line-height:1.85}
.toolbar{position:sticky;top:0;background:#143049;color:#fff;padding:12px 5vw;display:flex;justify-content:space-between;z-index:5;font-size:13px}.toolbar a{color:#c6f1ee;text-decoration:none}
main{max-width:1000px;margin:26px auto}.page{background:white;margin:0 0 26px;padding:60px 66px;box-shadow:0 8px 30px #1430490a;position:relative;scroll-margin-top:66px}
.eyebrow{font-size:11px;letter-spacing:2px;color:var(--teal);font-weight:700;text-transform:uppercase}.cover{min-height:960px;overflow:hidden;background:linear-gradient(135deg,#f9ffff 0%,#fff 55%,#e5f1f7 100%)}
.cover:after{content:'';position:absolute;border:45px solid #d5e9ed;border-radius:50%;width:280px;height:280px;right:-170px;top:40px;opacity:.55}
.brand{font-size:23px;font-weight:bold;letter-spacing:3px}.edition{color:var(--muted);font-size:12px;margin-top:8px}.cover h1{font-size:44px;line-height:1.45;letter-spacing:-1px;margin:90px 0 18px}.cover .subtitle{font-size:25px;color:var(--teal);margin:0 0 25px}.lead{font-size:16px;line-height:2;color:#466175;max-width:660px}.pillars{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin:54px 0 66px}.pillar{border-top:3px solid var(--teal);background:#eaf5f5;padding:21px 18px}.pillar b{display:block;font-size:17px}.pillar span{font-size:12px;color:var(--muted)}
.meta{font-size:11px;color:var(--muted)}h2{font-size:27px;line-height:1.5;margin:16px 0 25px;letter-spacing:-.4px}h3{font-size:17px;color:var(--teal);margin:25px 0 10px}p{margin:12px 0}table{width:100%;border-collapse:collapse;table-layout:fixed;margin:18px 0;font-size:13px;line-height:1.75}th{background:var(--ink);color:white;text-align:left;padding:13px 12px;font-weight:500}td{padding:13px 12px;border-bottom:1px solid var(--line);vertical-align:top;overflow-wrap:anywhere}tr:nth-child(even){background:#f0f6f8}th:first-child,td:first-child{width:24%}blockquote{margin:25px 0 0;padding:16px 20px;background:#eaf5f5;border-left:3px solid var(--teal);color:#0c6570;font-size:14px}.flow{display:flex;align-items:stretch;gap:8px;margin:20px 0;flex-wrap:wrap}.step{flex:1;min-width:70px;background:#edf5f7;padding:13px 8px;border-radius:4px;font-size:12px;text-align:center;font-weight:600}.arrow{align-self:center;color:var(--teal)}
.toc a{display:flex;padding:13px 0;border-bottom:1px solid var(--line);color:var(--ink);text-decoration:none;justify-content:space-between;font-size:15px}.toc a:hover{color:var(--teal)}ul{padding-left:23px}li{margin:6px 0}.pagefooter{border-top:1px solid var(--line);margin-top:30px;padding-top:10px;color:var(--muted);font-size:10px;display:flex;justify-content:space-between}
@media(max-width:720px){main{margin:0}.page{padding:32px 23px;margin-bottom:12px}.cover{min-height:850px}.cover h1{font-size:31px;margin-top:70px}.pillars{gap:8px}.pillar{padding:14px 10px}.pillar b{font-size:14px}h2{font-size:24px}table{font-size:12px}th,td{padding:10px 7px}}
@page{size:A4;margin:0}
@media print{html{scroll-behavior:auto}body{background:white;font-size:10.3pt;line-height:1.72;-webkit-print-color-adjust:exact;print-color-adjust:exact}.toolbar{display:none}main{max-width:none;margin:0}.page{width:210mm;min-height:297mm;margin:0;padding:17mm 18mm 15mm;box-shadow:none;break-after:page}.page:last-child{break-after:auto}.cover{height:297mm;min-height:0}.cover h1{font-size:33pt;margin-top:35mm}.cover .subtitle{font-size:22pt}.lead{font-size:12pt}.pillars{margin:22mm 0 26mm}.pillar{padding:6mm 4mm}.pillar b{font-size:12pt}.pillar span{font-size:9pt}.brand{font-size:20pt}.meta{font-size:8.5pt}.eyebrow{font-size:8pt}h2{font-size:21pt;margin:4mm 0 7mm}h3{font-size:12pt;margin:5mm 0 3mm}p{margin:3mm 0}table{font-size:9pt;margin:4mm 0;line-height:1.6}th,td{padding:2.7mm}tr,blockquote,.flow{break-inside:avoid}h2,h3{break-after:avoid}blockquote{font-size:9.5pt;padding:4mm 5mm;margin-top:6mm}.toc a{padding:4mm 0;font-size:11pt}.flow{gap:2mm}.step{font-size:8.5pt;min-width:14mm;padding:3mm 2mm}.pagefooter{font-size:7.5pt;margin-top:4mm}.page p{orphans:2;widows:2}#chapter-10{font-size:10pt;line-height:1.6}#chapter-10 li{margin:3px 0}#chapter-10 h3{margin-top:4mm}#chapter-10 blockquote{margin-top:4mm}}
"""


def render_body(body):
    out = []
    in_list = False
    for kind, value in blocks(body):
        if in_list and kind != "bullet":
            out.append("</ul>")
            in_list = False
        if kind == "table":
            out.append("<table><thead><tr>" + "".join(f"<th>{html.escape(c)}</th>" for c in value[0]) + "</tr></thead><tbody>")
            for row in value[1:]:
                out.append("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>")
            out.append("</tbody></table>")
        elif kind == "bullet":
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{html.escape(value)}</li>")
        elif kind == "flow":
            out.append('<div class="flow" aria-label="流程">' + '<span class="arrow">→</span>'.join(f'<span class="step">{html.escape(v)}</span>' for v in value.split(" → ")) + "</div>")
        else:
            tag = {"h3": "h3", "p": "p", "quote": "blockquote"}[kind]
            out.append(f"<{tag}>{html.escape(value)}</{tag}>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


def build_html(chapters):
    cover = '''<section class="page cover" id="cover"><div class="brand">KUNBUDDY / 坤伴</div><div class="edition">客户介绍手册 / 2026.09</div><h1>让需求推进更连贯<br>让交付依据更清晰</h1><p class="subtitle">Agent 协同研发平台</p><p class="lead">连接任务、项目事实、Agent 会话与代码交付，<br>让每一阶段都有具体产出，每一次交付都有核对依据。</p><div class="pillars"><div class="pillar"><b>协作更连贯</b><span>阶段产出有承接</span></div><div class="pillar"><b>过程可回看</b><span>状态与依据可核对</span></div><div class="pillar"><b>交付有控制</b><span>人工审批与版本记录</span></div></div><p class="meta">面向业务负责人、项目负责人、研发与信息化团队</p><p class="meta">版本日期：2026-09-18。本手册依据当前代码实现整理；具体开通范围、模型服务及实施安排以项目约定和部署验收为准。</p></section>'''
    toc = '<section class="page" id="contents"><div class="eyebrow">Reading guide</div><h2>阅读目录</h2><p>从平台价值开始，了解协作方式、交付机制与试点路径。</p><nav class="toc" aria-label="手册目录">'
    toc += "".join(f'<a href="#chapter-{i}"><span>{html.escape(title)}</span><span>↗</span></a>' for i, (title, _) in enumerate(chapters, 1))
    toc += '</nav><h3>建议阅读路径</h3><p>业务与决策负责人：01、02、09、10<br>项目与交付负责人：03、06、07、09<br>研发与信息化团队：04、05、06、08</p><p class="meta">文中流程和场景用于说明平台的使用方式，不代表客户案例或量化效果承诺。</p></section>'
    pages = [cover, toc]
    for i, (title, body) in enumerate(chapters, 1):
        pages.append(f'<section class="page" id="chapter-{i}"><div class="eyebrow">KUNBUDDY / Customer guide / {i:02d}</div><h2>{html.escape(title).replace("｜", "<br>")}</h2>{render_body(body)}<div class="pagefooter"><span>坤伴 · Agent 协同研发平台</span><span>2026.09 / {i:02d}</span></div></section>')
    output = '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta name="description" content="坤伴 Agent 协同研发平台客户介绍：协同效率、项目事实、交付治理与试点方式。"><title>坤伴 Agent 协同研发平台介绍手册</title><style>' + CSS + '</style></head><body><div class="toolbar"><span>KUNBUDDY / 客户介绍手册</span><a href="#contents">阅读目录</a></div><main>' + '\n'.join(pages) + '</main></body></html>'
    (OUT / (NAME + ".html")).write_text(output, encoding="utf-8")


def main():
    source = SOURCE.read_text(encoding="utf-8")
    parts = re.split(r"^## ", source, flags=re.M)[1:]
    chapters = [tuple(part.split("\n", 1)) for part in parts]
    assert len(chapters) == 10
    build_docx(chapters)
    build_html(chapters)
    print("Built DOCX and standalone HTML; 10 chapters; UTF-8 source.")


if __name__ == "__main__":
    main()
