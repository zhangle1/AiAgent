from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MANUAL = ROOT.parent / "kunbuddy-platform-manual-2026-09-10"
ASSETS = ROOT / "assets"

NAVY = "173B69"
BLUE = "2563EB"
PALE = "EEF5FF"
BORDER = "D9E2F0"
TEXT = "1F2937"
MUTED = "64748B"


def set_font(run, size: float, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), color)


def border(cell) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcBorders")
    if node is None:
        node = OxmlElement("w:tcBorders")
        props.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child = node.find(qn(f"w:{edge}"))
        if child is None:
            child = OxmlElement(f"w:{edge}")
            node.append(child)
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "6")
        child.set(qn("w:color"), BORDER)


def cell_margin(cell) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        props.append(node)
    for side, value in (("top", "120"), ("start", "140"), ("bottom", "120"), ("end", "140")):
        child = node.find(qn(f"w:{side}"))
        if child is None:
            child = OxmlElement(f"w:{side}")
            node.append(child)
        child.set(qn("w:w"), value)
        child.set(qn("w:type"), "dxa")


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    for name, size in (("Normal", 11), ("Title", 28), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name not in ("Normal", "Subtitle")
    title_props = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_props.find(qn("w:pBdr"))
    if title_border is not None:
        title_props.remove(title_border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("坤伴平台培训手册  |  2026-09-13")
    set_font(run, 8.5, False, MUTED)


def paragraph(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.45
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), 11, True)
        set_font(p.add_run(text[len(bold_prefix):]), 11)
    else:
        set_font(p.add_run(text), 11)


def heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(text), {1: 16, 2: 13, 3: 11.5}[level], True, "000000")


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = False
    for index, width in enumerate(widths):
        result.columns[index].width = Inches(width)
    for cell, text in zip(result.rows[0].cells, headers):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, NAVY)
        border(cell)
        cell_margin(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), 9.5, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        for cell, value in zip(result.add_row().cells, values):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            border(cell)
            cell_margin(cell)
            if row_index % 2:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def figure(doc: Document, image: Path, caption: str, width: float = 6.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    set_font(cap.add_run(caption), 9.5, False, MUTED)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), 10.5)


def copy_assets() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    source = MANUAL / "assets"
    mapping = {
        "login": source / "login-screen.png",
        "chat": source / "source-screenshots" / "screen-04.png",
        "project": source / "source-screenshots" / "screen-05.png",
        "model": source / "source-screenshots" / "screen-06.png",
        "answer": source / "source-screenshots" / "screen-07.png",
        "git": source / "source-screenshots" / "screen-08.png",
        "canvas": source / "source-screenshots" / "screen-13.png",
        "knowledge_flow": source / "knowledge-cps-aps-collaboration.png",
        "knowledge_roadmap": source / "knowledge-2.0-roadmap.png",
    }
    copied: dict[str, Path] = {}
    for key, path in mapping.items():
        target = ASSETS / f"{key}{path.suffix}"
        shutil.copy2(path, target)
        copied[key] = target
    return copied


def build_docx(images: dict[str, Path]) -> Path:
    doc = Document()
    configure(doc)
    cover = doc.add_paragraph(style="Title")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(72)
    cover.paragraph_format.space_after = Pt(12)
    set_font(cover.add_run("坤伴平台培训手册"), 28, True, "000000")
    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(28)
    set_font(sub.add_run("研发协作 代码交付 知识沉淀"), 13, False, MUTED)
    figure(doc, images["knowledge_flow"], "培训主题  受控事实驱动的 AI 协作", width=6.5)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ("培训对象：开发人员、交付人员、项目负责人、代码审批人", "建议时长：150 分钟  可拆分为两次 75 分钟培训", "更新日期：2026-09-13"):
        set_font(meta.add_run(line + "\n"), 10.5, False, MUTED)
    doc.add_page_break()

    heading(doc, "培训说明", 1)
    paragraph(doc, "本培训帮助学员在正确的项目与权限范围内使用坤伴完成查询、分析、协作、改动和受控交付。培训的重点不在于让 AI 自行决定，而在于让团队能够基于服务器代码、项目资料和 Git 记录形成可复核的工作结果。")
    table(doc, ["培训后应具备的能力", "最低完成标准"], [
        ["正确进入项目并选择上下文", "能说明当前项目、仓库和执行权限是否正确。"],
        ["提出可执行的问题", "能用目标、现状、限制、验收方式描述一次任务。"],
        ["审阅 AI 结果", "能核对引用的文件、Git 差异、验证记录和待确认项。"],
        ["完成受控交付", "知道变更集、审批、指纹复核和 Git 推送的顺序。"],
    ], [3.25, 3.85])
    heading(doc, "课程安排", 1)
    table(doc, ["时间", "模块", "方式", "学员产出"], [
        ["15 分钟", "平台边界与角色", "讲解与问答", "明确自己可操作的项目和权限。"],
        ["35 分钟", "登录 项目 会话上下文", "讲解加演示", "完成一次带约束的提问。"],
        ["30 分钟", "任务画布和协同", "案例演示", "能把复杂任务拆分为会话节点。"],
        ["40 分钟", "Git 变更集和交付", "上机练习", "完成模拟审阅和交付检查。"],
        ["30 分钟", "知识库和 CPS APS 协同", "规划讲解与讨论", "理解 2.0 建设方向和治理要求。"],
    ], [1.0, 2.3, 1.35, 2.45])

    heading(doc, "1 培训前准备", 1)
    table(doc, ["角色", "课前确认", "培训中的主要任务"], [
        ["开发人员", "已分配项目、已确认仓库可访问。", "分析代码、评审差异、验证修改。"],
        ["交付人员", "已准备需求说明、复现信息或验收标准。", "描述业务目标、追问原因、协同验收。"],
        ["项目负责人", "已确定本次练习项目和可用样例。", "组织会话、确认范围、安排审批。"],
        ["代码审批人", "已获得代码提交权限。", "审阅变更集并决定是否交付。"],
    ], [1.2, 2.75, 3.15])
    paragraph(doc, "培训环境请使用非生产或已授权的练习项目。禁止把密码、令牌、连接串、个人隐私和未脱敏生产数据放入聊天或上传附件。", bold_prefix="培训环境")
    heading(doc, "讲师课前检查", 2)
    table(doc, ["检查项", "确认方式"], [
        ["练习项目", "确认每位学员均能看见目标项目，并使用非生产或已授权代码库。"],
        ["演示账号", "准备开发人员和审批人两个角色的账号，避免在现场共享真实密码。"],
        ["样例任务", "准备一个只读分析题、一个差异审阅题和一个无需推送的交付演练题。"],
        ["网络和依赖", "确认浏览器可访问平台，相关模型、Git 和项目服务状态正常。"],
    ], [1.6, 5.5])
    doc.add_page_break()

    heading(doc, "2 第一次使用  登录和项目确认", 1)
    paragraph(doc, "第一步不是提问，而是确认登录账号、可见项目和代码库是否正确。项目决定可用的知识、文件、任务、模型和交付范围。")
    figure(doc, images["login"], "图 1  登录入口", width=5.9)
    figure(doc, images["chat"], "图 2  聊天工作台", width=6.35)
    figure(doc, images["project"], "图 3  在聊天中选择目标项目", width=6.35)
    heading(doc, "讲师演示", 2)
    for item in ("登录后检查左侧项目列表，确认目标项目存在。", "进入聊天，选择项目，再选择适合任务的 Agent、模型与执行权限。", "先用只读分析理解项目；需要改代码时，再由授权人员选择写入或完全控制。"):
        bullet(doc, item)

    heading(doc, "3 提问和结果审阅", 1)
    paragraph(doc, "高质量问题应包含四个部分：目标、现状、限制、验收方式。遇到 Bug 时，应提供症状、时间范围、相关模块、复现条件或日志片段，并要求结果列出代码和 Git 依据。")
    table(doc, ["练习类型", "可直接使用的提问模板", "应检查的证据"], [
        ["理解代码", "请说明 {模块} 的入口、核心调用链和涉及的配置文件；不要修改代码。", "文件路径、关键方法、引用关系、未知项。"],
        ["排查 Bug", "请分析 {现象} 在 {时间或版本} 出现的原因，并比较 {日期或提交} 前后的差异。", "相关提交、差异文件、可复现步骤、待验证假设。"],
        ["新功能规划", "基于当前项目事实，为 {目标} 形成实现计划，列出影响模块、风险、接口和验收条件。", "计划是否引用当前代码与需求，是否区分事实和建议。"],
        ["小型交付", "在 {项目} 中完成 {范围明确的改动}，先给出修改计划和验证方式，等待确认后再改。", "待提交文件、diff、测试结果、变更集状态。"],
    ], [1.15, 3.5, 2.45])
    figure(doc, images["model"], "图 4  选择模型和 Agent", width=6.35)
    figure(doc, images["answer"], "图 5  会话回答和实时过程", width=6.35)
    heading(doc, "上机练习", 2)
    paragraph(doc, "每名学员在练习项目中完成一次只读问题：选择一个模块，要求 AI 给出入口、依赖、风险和下一步验证动作。完成后由同组成员检查回答是否列出了可复核的文件或提交依据。")

    heading(doc, "4 任务画布和协同处理", 1)
    paragraph(doc, "简单任务可以在单个会话中处理。复杂任务应拆为分析、方案、改动、验证等会话节点，由项目负责人在画布上组织关系。新建会话默认不自动执行，避免把未确认的附件或需求直接转换为写入操作。")
    figure(doc, images["canvas"], "图 6  工作画布中的会话节点", width=6.55)
    table(doc, ["协同场景", "建议拆分", "交接时必须保留"], [
        ["交付提出需求，开发实现", "需求澄清 → 代码影响分析 → 实现 → 验证。", "需求版本、引用文件、验收条件、责任人。"],
        ["跨 CPS 和 APS 的影响分析", "APS 规则分析 → 接口影响复核 → CPS 改动建议 → 联合验证。", "使用的规则、接口约束、代码版本和结论。"],
        ["历史问题回溯", "现象收集 → Git 差异 → 原因假设 → 验证与修复。", "时间点、提交号、日志、验证结果。"],
    ], [1.8, 2.75, 2.55])

    heading(doc, "5 Git 交付演练", 1)
    paragraph(doc, "AI 可以帮助定位和修改，但正式交付仍由团队的权限与审批机制控制。学员需要知道哪些操作可以由自己完成，哪些必须交给代码审批人。")
    figure(doc, images["git"], "图 7  聊天中的 Git 状态和代码操作", width=6.35)
    table(doc, ["步骤", "学员操作", "检查点"], [
        ["1 确认范围", "确认当前项目、仓库、分支、需求与待提交文件。", "无无关文件；已保留必要的业务说明。"],
        ["2 审阅差异", "查看 diff、AI 的修改说明和验证结果。", "理解每个文件为什么改；不确定项标记为待确认。"],
        ["3 创建变更集", "形成修改摘要和校验记录，提交审批。", "变更集对应当前工作区内容。"],
        ["4 审批交付", "审批人复核范围、风险、测试和指纹。", "审批后若文件变化，需要重新审批。"],
        ["5 推送留痕", "获授权人员执行 Git 提交与推送。", "提交信息、远端状态和通知可追溯。"],
    ], [1.15, 3.25, 2.7])
    paragraph(doc, "练习规则：交付人员可以提出小型、范围明确的任务；开发人员负责核对代码事实与验证；审批人负责最终提交。Git 记录是团队复查和追踪的重要证据。", bold_prefix="练习规则")

    heading(doc, "6 知识库和 CPS APS Agent 协同展望", 1)
    paragraph(doc, "坤伴 2.0 将把代码、任务、需求文件、会话结论和交付记录沉淀为可检索、可引用、可审计的项目知识资产。CPS Agent 与 APS Agent 在同一项目权限边界内共享事实，而不是只交换自然语言结论。")
    figure(doc, images["knowledge_flow"], "图 8  知识库与 CPS APS Agent 协同流程", width=6.55)
    figure(doc, images["knowledge_roadmap"], "图 9  知识库建设路径", width=6.55)
    table(doc, ["阶段", "培训时应强调的原则"], [
        ["P0 数据接入", "资料进入知识库时必须记录来源、版本、权限和有效期。"],
        ["P1 知识可用", "回答应引用资源与版本；不确定结论必须说明待验证。"],
        ["P2 Agent 协同", "CPS 与 APS 的交接不得越过项目权限，且保留输入和结论。"],
        ["P3 自我演进", "高价值对话和技能候选要经过评估、灰度和回滚验证后才可复用。"],
    ], [1.45, 5.65])
    paragraph(doc, "参考方向：OpenViking 将资源、记忆和技能作为统一上下文并支持分层按需检索。坤伴会在自身项目权限、Git 交付和部署约束下形成适配方案，不应在培训中承诺尚未上线的接口或自动化能力。")

    heading(doc, "7 结课练习和检查表", 1)
    table(doc, ["练习", "学员应提交的结果", "讲师检查"], [
        ["项目确认", "截图或口述当前项目、仓库和执行模式。", "项目范围和权限选择是否正确。"],
        ["Bug 或差异问答", "一次带时间或提交范围的提问，以及 AI 给出的引用依据。", "是否包含文件、版本和待验证项。"],
        ["需求协同", "一个会话或画布拆分方案，包含输入资料和验收条件。", "是否区分需求资料与服务器代码事实。"],
        ["交付审阅", "待提交文件清单、diff 审阅结论和变更集操作说明。", "是否知道审批与 Git 留痕的顺序。"],
    ], [1.55, 3.4, 2.15])
    heading(doc, "常见问题", 2)
    for item in ("没有看到项目：先检查账号是否被分配项目，项目是否启用，再联系管理员。", "异常或未刷新：查看实时过程和后端日志，以当前工作区与 diff 为准；完全控制仅限受信任项目。"):
        bullet(doc, item)

    path = ROOT / "坤伴平台培训手册 2026-09-13.docx"
    doc.core_properties.title = "坤伴平台培训手册"
    doc.core_properties.subject = "研发协作与代码交付培训"
    doc.core_properties.author = "坤伴平台"
    doc.save(path)
    return path


def build_html(images: dict[str, Path]) -> Path:
    names = {key: path.name for key, path in images.items()}
    image = lambda key, caption: f'<figure><img src="assets/{names[key]}" alt="{caption}" loading="lazy"><figcaption>{caption}</figcaption></figure>'
    content = f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>坤伴平台培训手册</title><style>
:root{{--ink:#172033;--muted:#60708a;--line:#dbe5f2;--blue:#2563eb;--navy:#173b69;--soft:#eef5ff;--bg:#f6f8fc}}*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:var(--bg);color:var(--ink);font:16px/1.7 "Microsoft YaHei",Arial,sans-serif}}header{{background:linear-gradient(120deg,#173b69,#2563eb 60%,#6d28d9);color:#fff;padding:64px max(24px,calc((100vw - 1200px)/2))}}header h1{{margin:0;font-size:clamp(30px,4vw,50px)}}header p{{margin:10px 0 0;opacity:.9}}.layout{{max-width:1200px;margin:auto;display:grid;grid-template-columns:245px minmax(0,1fr);gap:28px;padding:28px 20px 64px}}nav{{position:sticky;top:16px;align-self:start;background:#fff;border:1px solid var(--line);border-radius:16px;padding:14px;box-shadow:0 8px 24px #18325a0d}}nav b{{display:block;padding:4px 9px 9px}}nav a{{display:block;padding:7px 9px;border-radius:8px;color:#334b68;text-decoration:none;font-size:14px}}nav a:hover{{background:var(--soft);color:var(--blue)}}section{{background:#fff;border:1px solid var(--line);border-radius:18px;padding:30px 32px;margin-bottom:20px;box-shadow:0 8px 24px #18325a0a}}h2{{margin:0 0 12px;font-size:26px;line-height:1.35}}h3{{margin:25px 0 8px;font-size:19px}}p{{margin:9px 0}}.tag{{font-size:12px;color:var(--blue);font-weight:700;letter-spacing:.12em}}table{{width:100%;border-collapse:collapse;margin:16px 0}}th,td{{border:1px solid var(--line);padding:10px 12px;text-align:left;vertical-align:top;font-size:14px}}th{{background:var(--navy);color:#fff}}tr:nth-child(even) td{{background:#f8fbff}}figure{{border:1px solid var(--line);border-radius:12px;overflow:hidden;margin:20px 0;background:#fff}}figure img{{display:block;width:100%;height:auto;cursor:zoom-in}}figcaption{{padding:8px 12px;color:var(--muted);font-size:13px}}.exercise{{padding:14px 16px;border:1px solid #bfdbfe;background:#f8fbff;border-radius:12px;margin-top:14px}}.exercise label{{display:block;padding:5px 0}}.toolbar{{display:flex;gap:10px;margin-top:18px}}input[type=search]{{width:100%;padding:10px;border:1px solid var(--line);border-radius:9px;font:inherit}}.hidden{{display:none!important}}#modal{{display:none;position:fixed;inset:0;z-index:9;align-items:center;justify-content:center;background:#071121cc;padding:20px}}#modal.show{{display:flex}}#modal img{{max-width:96vw;max-height:92vh;border-radius:8px}}@media(max-width:850px){{header{{padding:42px 20px}}.layout{{display:block;padding:16px 12px 44px}}nav{{position:relative;top:auto;display:flex;overflow:auto;white-space:nowrap;margin-bottom:16px}}nav b{{display:none}}nav a{{display:inline-block}}section{{padding:22px 16px}}table{{display:block;overflow:auto}}}}</style></head><body><header><div class="tag" style="color:#dbeafe">KUNBUDDY TRAINING</div><h1>坤伴平台培训手册</h1><p>围绕项目事实、AI 协作、受控改码与 Git 交付开展培训。</p></header><div class="layout"><nav><b>培训目录</b><a href="#intro">培训说明</a><a href="#prepare">课前准备</a><a href="#first">第一次使用</a><a href="#prompt">提问与审阅</a><a href="#canvas">任务与画布</a><a href="#delivery">Git 交付</a><a href="#knowledge">知识库展望</a><a href="#assessment">结课练习</a><input id="search" type="search" placeholder="搜索培训内容"></nav><main>
<section id="intro" class="searchable"><div class="tag">01</div><h2>培训说明</h2><p>培训的目标是让团队在正确的项目和权限边界内使用 AI，基于服务器代码、项目资料和 Git 记录形成可复核的结果。</p><table><thead><tr><th>模块</th><th>建议时长</th><th>学员产出</th></tr></thead><tbody><tr><td>平台边界与角色</td><td>15 分钟</td><td>确认自己的项目与权限。</td></tr><tr><td>登录 项目 上下文</td><td>35 分钟</td><td>一次带约束的提问。</td></tr><tr><td>任务画布和协同</td><td>30 分钟</td><td>复杂任务的会话拆分。</td></tr><tr><td>Git 变更集与交付</td><td>40 分钟</td><td>模拟审阅和交付检查。</td></tr><tr><td>知识库与多 Agent</td><td>30 分钟</td><td>理解建设方向与治理要求。</td></tr></tbody></table></section>
<section id="prepare" class="searchable"><div class="tag">02</div><h2>课前准备</h2><p>使用非生产或已授权的练习项目。不得把密码、令牌、连接串、隐私信息或未脱敏生产数据输入聊天或放入附件。</p><table><thead><tr><th>角色</th><th>课前确认</th><th>培训任务</th></tr></thead><tbody><tr><td>开发人员</td><td>项目和仓库可访问。</td><td>分析代码、评审差异、验证修改。</td></tr><tr><td>交付人员</td><td>准备需求、复现信息或验收标准。</td><td>描述业务目标、追问原因、协同验收。</td></tr><tr><td>项目负责人</td><td>确定练习项目和样例。</td><td>组织会话、确认范围、安排审批。</td></tr><tr><td>审批人</td><td>具备代码提交权限。</td><td>复核变更集，决定是否交付。</td></tr></tbody></table></section>
<section id="first" class="searchable"><div class="tag">03</div><h2>第一次使用</h2><p>先确认账号、可见项目和代码库，再开始提问。项目决定本次会话可用的知识、文件、任务、模型和交付范围。</p>{image('login','图 1 登录入口')}{image('chat','图 2 聊天工作台')}{image('project','图 3 选择目标项目')}<div class="exercise"><b>跟做清单</b><label><input type="checkbox" data-key="project">我确认了当前项目和仓库。</label><label><input type="checkbox" data-key="agent">我选择了适合任务的 Agent、模型和执行权限。</label><label><input type="checkbox" data-key="readonly">我知道理解项目应先用只读分析。</label></div></section>
<section id="prompt" class="searchable"><div class="tag">04</div><h2>提问和结果审阅</h2><p>任务描述应包含目标、现状、限制和验收方式。Bug 需补充症状、时间范围、相关模块、复现条件或日志片段，并要求返回代码与 Git 依据。</p><table><thead><tr><th>练习</th><th>模板</th><th>检查证据</th></tr></thead><tbody><tr><td>理解代码</td><td>说明模块入口、核心调用链和配置文件；不要修改代码。</td><td>文件路径、关键方法、未知项。</td></tr><tr><td>排查 Bug</td><td>分析现象在指定时间或版本出现的原因，并比较此前差异。</td><td>提交、差异文件、复现步骤、待验证假设。</td></tr><tr><td>新功能规划</td><td>列出影响模块、风险、接口和验收条件。</td><td>引用当前代码和需求，并区分事实与建议。</td></tr></tbody></table>{image('model','图 4 模型和 Agent 选择')}{image('answer','图 5 会话回答和实时过程')}<div class="exercise"><b>上机练习</b><label><input type="checkbox" data-key="question">我已完成一次只读问题。</label><label><input type="checkbox" data-key="evidence">我的回答包含了可复核的文件或提交依据。</label></div></section>
<section id="canvas" class="searchable"><div class="tag">05</div><h2>任务画布和协同</h2><p>简单任务可在单会话处理；复杂任务应拆为分析、方案、改动、验证等节点。新会话默认不自动运行，避免未确认资料直接触发写入。</p>{image('canvas','图 6 工作画布中的会话节点')}<table><thead><tr><th>场景</th><th>建议拆分</th><th>交接保留</th></tr></thead><tbody><tr><td>交付提出需求</td><td>澄清 → 影响分析 → 实现 → 验证。</td><td>需求版本、引用文件、验收条件、责任人。</td></tr><tr><td>CPS APS 影响分析</td><td>APS 规则 → 接口复核 → CPS 改动 → 联合验证。</td><td>规则、接口约束、代码版本、结论。</td></tr><tr><td>历史问题回溯</td><td>现象 → Git 差异 → 假设 → 验证修复。</td><td>时间点、提交号、日志、验证结果。</td></tr></tbody></table></section>
<section id="delivery" class="searchable"><div class="tag">06</div><h2>Git 交付演练</h2><p>AI 可帮助定位和修改，正式交付仍由团队的权限与审批机制控制。</p>{image('git','图 7 聊天中的 Git 状态和代码操作')}<table><thead><tr><th>步骤</th><th>操作</th><th>检查点</th></tr></thead><tbody><tr><td>确认范围</td><td>确认项目、仓库、分支和待提交文件。</td><td>无无关文件，业务说明齐全。</td></tr><tr><td>审阅差异</td><td>查看 diff、修改说明和验证结果。</td><td>理解每个文件为何修改。</td></tr><tr><td>创建变更集</td><td>形成摘要和校验记录，提交审批。</td><td>内容与工作区一致。</td></tr><tr><td>审批交付</td><td>复核范围、风险、测试和指纹。</td><td>审批后变化需重新审批。</td></tr><tr><td>推送留痕</td><td>获授权者提交并推送。</td><td>提交信息与远端状态可追溯。</td></tr></tbody></table><div class="exercise"><b>练习规则</b><label><input type="checkbox" data-key="diff">我已审阅待提交文件和 diff。</label><label><input type="checkbox" data-key="approval">我知道提交前需要变更集和审批。</label><label><input type="checkbox" data-key="trace">我知道 Git 记录是证据链的一部分。</label></div></section>
<section id="knowledge" class="searchable"><div class="tag">07</div><h2>知识库和 CPS APS Agent 协同展望</h2><p>坤伴 2.0 计划将受控项目的代码、任务、需求文件、会话结论和交付记录沉淀为项目知识资产，CPS Agent 与 APS Agent 在同一权限边界内共享事实。</p>{image('knowledge_flow','图 8 知识库与 CPS APS Agent 协同流程')}{image('knowledge_roadmap','图 9 知识库建设路径')}<p>参考 OpenViking 的资源、记忆和技能统一组织，以及摘要、概览、明细的分层取用思路。坤伴会结合项目权限、Git 交付和部署约束建设自身适配层；培训中不应承诺尚未上线的接口或自动化能力。</p></section>
<section id="assessment" class="searchable"><div class="tag">08</div><h2>结课练习和检查</h2><table><thead><tr><th>练习</th><th>提交结果</th><th>讲师检查</th></tr></thead><tbody><tr><td>项目确认</td><td>当前项目、仓库和执行模式。</td><td>范围和权限正确。</td></tr><tr><td>Bug 或差异问答</td><td>含时间或提交范围的提问及引用依据。</td><td>文件、版本和待验证项完整。</td></tr><tr><td>需求协同</td><td>会话或画布拆分方案。</td><td>区分需求资料和代码事实。</td></tr><tr><td>交付审阅</td><td>文件清单、diff 结论和变更集说明。</td><td>理解审批与 Git 留痕顺序。</td></tr></tbody></table><div class="exercise"><b>我已完成</b><label><input type="checkbox" data-key="finish1">一次受限范围的 AI 问答。</label><label><input type="checkbox" data-key="finish2">一次证据化结果审阅。</label><label><input type="checkbox" data-key="finish3">一次受控交付流程演练。</label></div></section></main></div><div id="modal"><img alt="图片预览"></div><script>const modal=document.querySelector('#modal'),mi=modal.querySelector('img');document.querySelectorAll('figure img').forEach(i=>i.onclick=()=>{{mi.src=i.src;modal.classList.add('show')}});modal.onclick=()=>modal.classList.remove('show');document.querySelectorAll('input[type=checkbox]').forEach(i=>{{i.checked=localStorage.getItem('kun-training-'+i.dataset.key)==='1';i.onchange=()=>localStorage.setItem('kun-training-'+i.dataset.key,i.checked?'1':'0')}});document.querySelector('#search').oninput=e=>{{const q=e.target.value.toLowerCase();document.querySelectorAll('.searchable').forEach(s=>s.classList.toggle('hidden',q&&!s.innerText.toLowerCase().includes(q)))}};</script></body></html>'''
    path = ROOT / "坤伴平台培训手册 2026-09-13.html"
    path.write_text(content, encoding="utf-8")
    return path


def main() -> None:
    images = copy_assets()
    docx = build_docx(images)
    html = build_html(images)
    print(f"{{'docx': '{docx}', 'html': '{html}'}}")


if __name__ == "__main__":
    main()
