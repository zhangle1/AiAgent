"""Build the two-page customer brief: python build_manual_assets.py.

Requires python-docx, playwright (with Chromium), and pymupdf.
The supplied screenshot is preserved as-is; all formats share this content.
"""
from pathlib import Path
from html import escape
from zipfile import ZipFile, ZIP_DEFLATED
import os

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import sync_playwright
import pymupdf

ROOT = Path(__file__).resolve().parent
STEM = "坤伴Agent协同研发平台介绍手册"
TITLE = "让 AI 融入项目，让协作走向交付"
INTRO = "坤伴面向企业研发团队，将 AI 会话、项目代码、知识资料与研发工具汇聚到同一工作台，帮助团队从理解需求、形成方案，推进到开发验证与审批交付。"
GUIDE = [
    ("选项目", "按项目组织会话与代码，延续任务上下文。"),
    ("说需求", "结合文字、资料与图片，开展分析和实施。"),
    ("看结果", "查看代码差异，使用运行、打包与交付入口。"),
]
VALUE = [
    ("沟通更聚焦", "围绕项目事实讨论，减少重复说明。"),
    ("成果更直观", "通过原型、预览和代码差异核对结果。"),
    ("经验可复用", "保留资料来源，让知识服务后续任务。"),
]
FEATURES = [
    ("项目会话", "关联项目与代码上下文，辅助理解系统、分析问题和修改代码。"),
    ("知识中心", "保留原文，按需提炼知识草稿，支持检索与来源核对。"),
    ("工作画布", "编排多个 Agent 会话的职责和上下游关系，查看执行状态。"),
    ("原型与看板", "将需求转成可预览的页面，支持评审、修改与效果确认。"),
    ("验证与交付", "查看差异、运行与打包；通过变更集检查、人工审批后交付。"),
    ("持续养护", "分析存量代码、形成改进建议，支持验证后生成待审批变更集。"),
]
STEPS = [
    ("明确需求", "目标 · 资料 · 验收条件"),
    ("协同实施", "分析 · 原型 · 代码修改"),
    ("验证交付", "检查 · 审批 · 版本记录"),
]
CTA = "从一个真实需求开始"
CTA_BODY = "选择一个项目、一项明确任务和一组验收条件，用可预览的成果与可核对的交付记录，评估协作效果。"
NOTE = "支持企业内网部署；模型与外部连接按项目配置。具体能力以部署版本、权限和验收范围为准。"


def write_text(path, text):
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(text, encoding="utf-8")
    os.replace(temp, path)


def cards(items, cls="card"):
    return "".join(f'<div class="{cls}"><h3>{escape(a)}</h3><p>{escape(b)}</p></div>' for a, b in items)


def build_html():
    css = """*{box-sizing:border-box}body{margin:0;background:#edf2f9;color:#182844;font:14px/1.65,'Microsoft YaHei','Noto Sans CJK SC',sans-serif}
.page{width:210mm;height:297mm;margin:24px auto;padding:17mm 17mm 15mm;background:white;position:relative;break-after:page}
.page:last-child{break-after:auto}.brand{display:flex;align-items:center;justify-content:space-between;color:#285ed9;font-weight:700;font-size:17px;border-bottom:1px solid #dfe7f3;padding-bottom:13px}
.brand span{font-size:10px;letter-spacing:2px;font-weight:500;color:#76869f}.eyebrow{color:#2860df;font-size:11px;letter-spacing:2px;margin:25px 0 10px}
h1{font-size:34px;line-height:1.4;letter-spacing:-1px;margin:0 0 14px}h2{font-size:28px;line-height:1.4;margin:0 0 12px}h3{font-size:15px;margin:0 0 7px}p{margin:0;color:#596980}.intro{font-size:14px;line-height:1.9;max-width:630px}
.screen{margin:23px 0 0;padding:9px;background:#f4f7fc;border:1px solid #dfe7f3;border-radius:12px}.screen img{display:block;width:100%;border-radius:6px}.caption{font-size:10px;color:#8090a6;margin:8px 0 18px}
.guide{display:grid;grid-template-columns:repeat(3,1fr);gap:16px;margin:0 0 25px}.guide h3{color:#2860df}.guide p{font-size:12px}.section-label{font-size:11px;color:#71829c;letter-spacing:2px;margin:0 0 12px}
.values{background:#f1f5ff;border-radius:12px;padding:20px;display:grid;grid-template-columns:repeat(3,1fr);gap:16px}.values p{font-size:12px}.quote{font-size:19px;color:#214c9b;margin-top:25px;font-weight:600;line-height:1.7}
.features{display:grid;grid-template-columns:1fr 1fr;gap:13px;margin:22px 0}.card{border:1px solid #dfe7f3;border-radius:10px;padding:16px 18px;min-height:112px}.card h3{color:#204fae}.card p{font-size:12px;line-height:1.85}.flow{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin:12px 0 22px}.step{border-top:3px solid #2c63e5;padding-top:12px}.step p{font-size:11px}.cta{padding:23px;background:#235be0;color:white;border-radius:12px}.cta h3{font-size:22px;margin-bottom:8px}.cta p{color:#eef4ff;font-size:13px;line-height:1.9}.note{font-size:10px;line-height:1.8;margin-top:16px}.footer{position:absolute;bottom:12mm;left:17mm;right:17mm;display:flex;justify-content:space-between;border-top:1px solid #dfe7f3;padding-top:10px;color:#8491a5;font-size:10px}
@page{size:A4;margin:0}@media print{body{background:white}.page{margin:0;print-color-adjust:exact;-webkit-print-color-adjust:exact}}
@media screen and (max-width:800px){.page{width:100%;height:auto;min-height:100vh;padding:24px;margin:0 0 16px}.footer{position:static;margin-top:28px}h1{font-size:29px}.guide,.values,.flow{grid-template-columns:1fr}.features{grid-template-columns:1fr}}
""".replace("\\+", "")
    brand = '<header class="brand">坤伴 <span>AGENT 协同研发平台</span></header>'
    def footer(n):
        return f'<footer class="footer"><span>坤伴 · 产品介绍 / 2026.09</span><span>0{n} / 02</span></footer>'
    html = f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>坤伴 Agent｜两页产品介绍</title><style>{css}</style></head><body>
<section class="page">{brand}<div class="eyebrow">项目 · 知识 · 研发协作</div><h1>让 AI 融入项目<br>让协作走向交付</h1><p class="intro">{INTRO}</p>
<figure class="screen"><img src="assets/main-workspace.png" alt="坤伴主界面：左侧项目会话、中间 AI 对话、右侧项目程序运行面板"></figure><p class="caption">产品主界面实拍 · 项目会话、AI 对话与项目程序运行集中呈现</p>
<div class="guide">{cards(GUIDE, 'guide-item')}</div><div class="section-label">为团队带来的价值</div><div class="values">{cards(VALUE, 'value')}</div><div class="quote">把需求说清楚，把成果做出来，<br>让每一步都有依据。</div>{footer(1)}</section>
<section class="page">{brand}<div class="eyebrow">核心能力 / 工作方式</div><h2>从需求到交付，连贯推进</h2><p class="intro">连接项目上下文与研发工具，让分析、实施和复核围绕同一项任务展开。</p><div class="features">{cards(FEATURES)}</div>
<div class="section-label">一条清晰的协作路径</div><div class="flow">{cards(STEPS, 'step')}</div><div class="cta"><h3>{CTA}</h3><p>{CTA_BODY}</p></div><p class="note">{NOTE}</p>{footer(2)}</section></body></html>'''
    write_text(ROOT / f"{STEM}.html", html)


def build_markdown():
    text = f"# 坤伴 Agent 协同研发平台\n\n## 第 1 页｜{TITLE}\n\n{INTRO}\n\n![坤伴主界面](assets/main-workspace.png)\n\n"
    for title, items in [("主界面，一站协作", GUIDE), ("为团队带来的价值", VALUE)]:
        text += f"### {title}\n\n" + "\n".join(f"- **{a}**：{b}" for a,b in items) + "\n\n"
    text += "## 第 2 页｜从需求到交付，连贯推进\n\n" + "\n".join(f"- **{a}**：{b}" for a,b in FEATURES)
    text += "\n\n### 一条清晰的协作路径\n\n" + " → ".join(a for a,b in STEPS)
    text += f"\n\n### {CTA}\n\n{CTA_BODY}\n\n{NOTE}\n\n2026 年 9 月 · 两页精简版\n"
    write_text(ROOT / f"{STEM}.md", text)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(17)
    sec.left_margin = sec.right_margin = Mm(18)
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.3
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        st = doc.styles[name]
        st.font.name = 'Microsoft YaHei'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.color.rgb = RGBColor.from_string('235BE0')
    sec.header.paragraphs[0].text = '坤伴  /  AGENT 协同研发平台'
    foot = sec.footer.paragraphs[0]
    foot.add_run('坤伴 · 产品介绍 / 2026.09                                         ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    foot.add_run(' / 2')
    doc.add_heading('让 AI 融入项目\n让协作走向交付', 0)
    doc.add_paragraph(INTRO)
    doc.add_picture(str(ROOT / 'assets/main-workspace.png'), width=Mm(174))
    p = doc.add_paragraph('产品主界面实拍 · 项目会话、AI 对话与项目程序运行集中呈现')
    p.runs[0].font.size = Pt(8)
    def item(a,b):
        p = doc.add_paragraph(); p.add_run(a + '  ').bold = True; p.add_run(b)
    doc.add_heading('主界面，一站协作', 2)
    for a,b in GUIDE: item(a,b)
    doc.add_heading('为团队带来的价值', 2)
    for a,b in VALUE: item(a,b)
    doc.add_paragraph('把需求说清楚，把成果做出来，让每一步都有依据。')
    doc.add_page_break()
    doc.add_heading('从需求到交付，连贯推进', 0)
    doc.add_paragraph('连接项目上下文与研发工具，让分析、实施和复核围绕同一项任务展开。')
    for a,b in FEATURES:
        doc.add_heading(a, 3); doc.add_paragraph(b)
    doc.add_heading('一条清晰的协作路径', 2)
    doc.add_paragraph('  →  '.join(a for a,b in STEPS))
    doc.add_heading(CTA, 2); doc.add_paragraph(CTA_BODY)
    doc.add_paragraph(NOTE)
    temp = ROOT / f'{STEM}.tmp.docx'; doc.save(temp); os.replace(temp, ROOT / f'{STEM}.docx')


def render_pdf():
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1100, "height": 1200})
        page.goto((ROOT / f'{STEM}.html').as_uri(), wait_until='networkidle')
        page.emulate_media(media='print')
        page.evaluate('document.fonts.ready')
        assert page.locator('img').evaluate_all('(imgs) => imgs.every(i => i.complete && i.naturalWidth > 0)')
        assert page.locator('.page').evaluate_all('(ps) => ps.every(p => p.scrollHeight <= p.clientHeight)'), 'Page overflow'
        temp = ROOT / f'{STEM}.tmp.pdf'
        page.pdf(path=str(temp), print_background=True, prefer_css_page_size=True)
        browser.close()
    with pymupdf.open(temp) as pdf:
        assert len(pdf) == 2, f'Expected 2 pages, got {len(pdf)}'
        assert all(len(p.get_text()) > 100 for p in pdf), 'Missing text'
    os.replace(temp, ROOT / f'{STEM}.pdf')


def package():
    temp = ROOT / '坤伴介绍.tmp.zip'
    with ZipFile(temp, 'w', ZIP_DEFLATED) as archive:
        for ext in ['pdf','docx','html','md']:
            archive.write(ROOT / f'{STEM}.{ext}', f'{STEM}.{ext}')
        archive.write(ROOT / 'assets/main-workspace.png', 'assets/main-workspace.png')
    os.replace(temp, ROOT / '坤伴介绍.zip')


if __name__ == '__main__':
    build_html(); build_markdown(); build_docx(); render_pdf(); package()
    print('Built HTML, Markdown, Word, PDF (verified: 2 pages), and ZIP.')
